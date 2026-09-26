#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script to generate the comprehensive quotation and functional breakdown Excel file for DuaxCar Kitchen
Matching the structure and styling of:
/Users/tungnguyen/Tulie/SeaWay/Bao_gia_Website_Gia_vi_Sot_v6_Pham_vi_cu_the.xlsx
"""

import os
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_duaxcar_quotation():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    font_family = "Arial"
    
    # Palette
    c_navy = "17365D"       # Main header dark navy
    c_white = "FFFFFF"
    c_gold = "C79A2B"       # Parameters header
    c_gray_bg = "E7E6E6"    # Core background header
    c_green_bg = "E2F0D9"   # Full background header
    c_green_txt = "2E7D32"  # Full dark green text
    c_core_txt = "17365D"   # Core dark navy text
    c_yellow_accent = "FFFFF8DC" # Accent yellow for prices
    c_sub_yellow = "FFFFF2CC"
    c_check_bg = "EEF3F8"   # Light blue tint for checkmarks
    c_alt_row = "F8FAFC"    # Alternating row background
    c_border = "D9D9D9"     # Light gray border

    thin_border_side = Side(border_style="thin", color=c_border)
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=Side(border_style="medium", color=c_navy))

    # =========================================================================
    # SHEET 1: 01_BAO_GIA_TONG_HOP
    # =========================================================================
    ws1 = wb.create_sheet(title="01_BAO_GIA_TONG_HOP")
    ws1.views.sheetView[0].showGridLines = True

    # Title
    ws1.merge_cells("A1:I1")
    ws1["A1"] = "ĐỀ XUẤT WEBSITE HỌC VIỆN ẨM THỰC DUAXCAR KITCHEN – ĐÀO TẠO NGHỀ & CỐ VẤN KINH DOANH F&B"
    ws1["A1"].font = Font(name=font_family, size=16, bold=True, color=c_white)
    ws1["A1"].fill = PatternFill(start_color=c_navy, end_color=c_navy, fill_type="solid")
    ws1["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws1.row_dimensions[1].height = 42

    ws1.merge_cells("A2:I2")
    ws1["A2"] = "Phạm vi gồm 02 phương án: Gói Tiết kiệm (Core Academy) và Gói Full (Học viện & Chuyển giao F&B Toàn diện)."
    ws1["A2"].font = Font(name=font_family, size=10, italic=True, color="595959")
    ws1["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws1.row_dimensions[2].height = 24

    # Parameter Box
    # Row 4 Headers
    ws1["A4"] = "THAM SỐ"
    ws1["B4"] = "Giá trị"
    ws1.merge_cells("D4:E4")
    ws1["D4"] = "GÓI TIẾT KIỆM – CORE ACADEMY"
    ws1.merge_cells("G4:H4")
    ws1["G4"] = "GÓI FULL – KHUYẾN NGHỊ TĂNG TRƯỞNG"

    for c in ["A4", "B4"]:
        ws1[c].font = Font(name=font_family, size=10, bold=True, color=c_white)
        ws1[c].fill = PatternFill(start_color=c_gold, end_color=c_gold, fill_type="solid")
        ws1[c].alignment = Alignment(horizontal="center", vertical="center")
        ws1[c].border = cell_border

    ws1["D4"].font = Font(name=font_family, size=10, bold=True, color=c_core_txt)
    ws1["D4"].fill = PatternFill(start_color=c_gray_bg, end_color=c_gray_bg, fill_type="solid")
    ws1["D4"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["D4"].border = cell_border
    ws1["E4"].border = cell_border

    ws1["G4"].font = Font(name=font_family, size=10, bold=True, color=c_green_txt)
    ws1["G4"].fill = PatternFill(start_color=c_green_bg, end_color=c_green_bg, fill_type="solid")
    ws1["G4"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["G4"].border = cell_border
    ws1["H4"].border = cell_border
    ws1.row_dimensions[4].height = 24

    # Row 5: Discount & Module count
    ws1["A5"] = "Chiết khấu Gói Tiết kiệm"
    ws1["B5"] = 0.20
    ws1["D5"] = "Module quy trình"
    ws1["E5"] = 10
    ws1["G5"] = "Module quy trình"
    ws1["H5"] = 16

    # Row 6: Features
    ws1["A6"] = "Chiết khấu Gói Full"
    ws1["B6"] = 0.20
    ws1["D6"] = "Feature/hạng mục"
    ws1["E6"] = "='02_CHI_TIET_MODULE'!F324"  # Will point to calculated feature count
    ws1["G6"] = "Feature/hạng mục"
    ws1["H6"] = "='02_CHI_TIET_MODULE'!G324"

    # Row 7: Screens
    ws1["D7"] = "Trang/UI"
    ws1["E7"] = 29
    ws1["G7"] = "Trang/UI"
    ws1["H7"] = 40

    # Row 8: List price
    ws1.merge_cells("A8:B8")
    ws1["A8"] = "CHÊNH LỆCH FULL"
    ws1["D8"] = "Giá niêm yết"
    ws1["E8"] = "=SUMIF($F$14:$F$29,\"✓\",$H$14:$H$29)"
    ws1["G8"] = "Giá niêm yết"
    ws1["H8"] = "=SUMIF($G$14:$G$29,\"✓\",$H$14:$H$29)"

    # Row 9: Total package
    ws1["A9"] = "Tăng thêm"
    ws1["B9"] = "=H9-E9"
    ws1["D9"] = "TỔNG GÓI"
    ws1["E9"] = "=E8*(1-$B$5)"
    ws1["G9"] = "TỔNG GÓI"
    ws1["H9"] = "=H8*(1-$B$6)"

    # Row 10: Ratio increase
    ws1["A10"] = "Tỷ lệ tăng"
    ws1["B10"] = "=IF(E9=0,0,B9/E9)"

    # Row 11: Add-on Full
    ws1["A11"] = "Add-on Full"
    ws1["B11"] = 6

    # Formatting parameter box
    for r in range(5, 12):
        ws1.row_dimensions[r].height = 20
        for col_idx in [1, 2, 4, 5, 7, 8]:
            c_letter = get_column_letter(col_idx)
            cell = ws1[f"{c_letter}{r}"]
            cell.font = Font(name=font_family, size=10)
            cell.border = cell_border
            if col_idx in [1, 4, 7]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")

    ws1["B5"].number_format = "0%"
    ws1["B5"].fill = PatternFill(start_color=c_yellow_accent, end_color=c_yellow_accent, fill_type="solid")
    ws1["B6"].number_format = "0%"
    ws1["B6"].fill = PatternFill(start_color=c_yellow_accent, end_color=c_yellow_accent, fill_type="solid")
    ws1["B9"].number_format = '#,##0\\ "₫"'
    ws1["B9"].fill = PatternFill(start_color=c_sub_yellow, end_color=c_sub_yellow, fill_type="solid")
    ws1["B10"].number_format = "0.0%"

    ws1["A8"].font = Font(name=font_family, size=10, bold=True, color=c_white)
    ws1["A8"].fill = PatternFill(start_color=c_navy, end_color=c_navy, fill_type="solid")
    ws1["A8"].alignment = Alignment(horizontal="center", vertical="center")

    ws1["D5"].font = Font(name=font_family, size=10, bold=True, color=c_core_txt)
    ws1["D6"].font = Font(name=font_family, size=10, bold=True, color=c_core_txt)
    ws1["D7"].font = Font(name=font_family, size=10, bold=True, color=c_core_txt)
    ws1["D8"].font = Font(name=font_family, size=10, bold=True, color=c_core_txt)
    ws1["E8"].font = Font(name=font_family, size=10, bold=True)
    ws1["E8"].number_format = '#,##0\\ "₫"'

    ws1["G5"].font = Font(name=font_family, size=10, bold=True, color=c_green_txt)
    ws1["G6"].font = Font(name=font_family, size=10, bold=True, color=c_green_txt)
    ws1["G7"].font = Font(name=font_family, size=10, bold=True, color=c_green_txt)
    ws1["G8"].font = Font(name=font_family, size=10, bold=True, color=c_green_txt)
    ws1["H8"].font = Font(name=font_family, size=10, bold=True)
    ws1["H8"].number_format = '#,##0\\ "₫"'

    # Total package row styling
    ws1["D9"].font = Font(name=font_family, size=10, bold=True, color=c_core_txt)
    ws1["D9"].fill = PatternFill(start_color=c_gray_bg, end_color=c_gray_bg, fill_type="solid")
    ws1["E9"].font = Font(name=font_family, size=10, bold=True, color=c_core_txt)
    ws1["E9"].fill = PatternFill(start_color=c_gray_bg, end_color=c_gray_bg, fill_type="solid")
    ws1["E9"].number_format = '#,##0\\ "₫"'

    ws1["G9"].font = Font(name=font_family, size=10, bold=True, color=c_green_txt)
    ws1["G9"].fill = PatternFill(start_color=c_green_bg, end_color=c_green_bg, fill_type="solid")
    ws1["H9"].font = Font(name=font_family, size=10, bold=True, color=c_green_txt)
    ws1["H9"].fill = PatternFill(start_color=c_green_bg, end_color=c_green_bg, fill_type="solid")
    ws1["H9"].number_format = '#,##0\\ "₫"'

    # Master Table Header (Row 13)
    headers_ws1 = ["Mã", "Nhóm", "MODULE QUY TRÌNH", "Kết quả / Giá trị bàn giao", "Số hạng mục", "Tiết kiệm", "Full", "Giá module tham chiếu", "Ghi chú"]
    for c_idx, h in enumerate(headers_ws1, 1):
        cell = ws1.cell(13, c_idx, h)
        cell.font = Font(name=font_family, size=10, bold=True, color=c_white)
        cell.fill = PatternFill(start_color=c_navy, end_color=c_navy, fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if c_idx in [1, 5, 6, 7] else "left", vertical="center")
        cell.border = cell_border
    ws1.row_dimensions[13].height = 26

    # Modules Data for DuaxCar Kitchen
    modules_data = [
        ("M01", "NỀN TẢNG", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "Giao diện thương hiệu chuẩn nhận diện DuaxCar (Cam-Đen), mobile-first, tạo dựng uy tín đào tạo ẩm thực chuyên nghiệp.", 24, "✓", "✓", 10000000, "CORE – Không tách lẻ"),
        ("M02", "NỘI DUNG", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "Quản trị bài viết, cơ sở vật chất bếp thực tế, FAQ, các văn bản chính sách đào tạo, quy chế học viên và bảo lưu.", 24, "✓", "✓", 5000000, "CORE – Không tách lẻ"),
        ("M03", "ĐÀO TẠO", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "Học viên tra cứu khóa học theo hình thức (Onsite, 1-kèm-1, Nghề mở quán, E-learning) và theo chuyên đề món ăn.", 23, "✓", "✓", 5000000, "CORE – Không tách lẻ"),
        ("M04", "CURRICULUM", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "Trình bày giáo trình theo buổi học (Curriculum Accordion), bí quyết nêm sốt cốt, hỗ trợ tính cost và setup mở quán.", 26, "✓", "✓", 10000000, "CORE – Không tách lẻ"),
        ("M05", "GIẢNG VIÊN", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "Hồ sơ chuyên môn, danh hiệu Nghệ nhân Bàn tay vàng, 25+ năm kinh nghiệm bếp trưởng, bảo chứng chất lượng đào tạo.", 16, "✓", "✓", 5000000, "CORE – Không tách lẻ"),
        ("M06", "LỊCH HỌC", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "Cập nhật lịch khai giảng theo tháng, ca học sáng/chiều/tối, trạng thái lớp (còn chỗ/sắp đủ/đã chốt) và địa điểm học.", 16, "✓", "✓", 5000000, "CORE – Không tách lẻ"),
        ("M07", "TUYỂN SINH", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "Form & Modal đăng ký tư vấn giữ chỗ thông minh, chuẩn hóa số điện thoại, phân loại nhu cầu học viên và chống spam.", 20, "✓", "✓", 5000000, "CORE – Không tách lẻ"),
        ("M08", "VẬN HÀNH", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "Bảng điều khiển CMS quản lý lead tuyển sinh, cập nhật trạng thái tư vấn, chỉnh sửa khóa học, giáo trình và giảng viên.", 22, "✓", "✓", 10000000, "CORE – Không tách lẻ"),
        ("M09", "MEDIA", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "Quản lý hình ảnh món ăn, video bếp; tự động nén WebP trước khi tải lên, tích hợp Media Selector dùng chung.", 16, "✓", "✓", 5000000, "CORE – Không tách lẻ"),
        ("M10", "CONTENT", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "Chia sẻ bí quyết nghề bếp, kinh nghiệm quản trị F&B, mục lục tự động (Table of Contents), trình soạn thảo Rich Text.", 22, "✓", "✓", 10000000, "CORE – Không tách lẻ"),
        ("M11", "MARKETING", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Tracking", "Course Schema, EducationalOrganization Schema, Sitemap tự động, tích hợp GA4, GTM và Meta Pixel đo lường tuyển sinh.", 20, "—", "✓", 5000000, "ADD-ON – Nâng Cao Hiệu Quả Tuyển Sinh"),
        ("M12", "KỸ THUẬT", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "Hạ tầng Vercel Edge Serverless, Supabase PostgreSQL RLS, Cron Keep-alive 24/7, tự động backup dữ liệu định kỳ.", 20, "—", "✓", 10000000, "ADD-ON – Vận Hành Ổn Định & An Toàn"),
        ("M13", "THANH TOÁN", "Cổng Đặt cọc Giữ chỗ Khóa học & Học phí Trực tuyến", "Tự sinh mã VietQR động đúng số tiền cọc/học phí và mã đăng ký; tích hợp cổng VNPay/MoMo và webhook đối soát tự động.", 18, "—", "✓", 10000000, "ADD-ON – Tự Động Hóa Dòng Tiền Tuyển Sinh"),
        ("M14", "TRUST & HẬU MÃI", "Social Proof Chuyên sâu & Học viên Mở quán Thành công", "Showcase quán ăn học viên đã mở, phỏng vấn thực tế, review xác thực kèm ảnh thành phẩm và hệ thống kiểm duyệt.", 18, "—", "✓", 10000000, "ADD-ON – Tăng Trưởng Uy Tín Chuyển Đổi"),
        ("M15", "HỌC VIÊN", "Cổng Học viên (Student Portal) & Kho Công thức Số", "Tài khoản học viên xem lịch học, tải tài liệu công thức điện tử, tự động đóng Watermark SĐT/tên chống phát tán công thức.", 16, "—", "✓", 10000000, "ADD-ON – Giữ Chân & Bảo Vệ Bản Quyền F&B"),
        ("M16", "AUTOMATION", "CRM Tuyển sinh Tự động, Email / ZNS & Vận hành Đa cấp", "Tự động gửi email/ZNS xác nhận & nhắc lịch học, webhook đồng bộ CRM ngoài, phân quyền nhân viên và báo cáo doanh thu.", 18, "—", "✓", 5000000, "ADD-ON – Scale Quy Mô Học Viện")
    ]

    for idx, row_data in enumerate(modules_data, 14):
        ws1.row_dimensions[idx].height = 24
        m_code, m_group, m_name, m_val, m_count, m_core, m_full, m_price, m_note = row_data
        
        ws1.cell(idx, 1, m_code).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(idx, 1).font = Font(name=font_family, size=10, bold=True, color=c_core_txt)
        
        ws1.cell(idx, 2, m_group).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(idx, 2).font = Font(name=font_family, size=10, bold=True, color=c_core_txt)
        
        ws1.cell(idx, 3, m_name).alignment = Alignment(horizontal="left", vertical="center")
        ws1.cell(idx, 3).font = Font(name=font_family, size=10, bold=True, color=c_core_txt)
        
        ws1.cell(idx, 4, m_val).alignment = Alignment(horizontal="left", vertical="center")
        ws1.cell(idx, 4).font = Font(name=font_family, size=9)
        
        # Link item count from sheet 2 dynamically via formula:
        ws1.cell(idx, 5, f"=COUNTIF('02_CHI_TIET_MODULE'!$A$5:$A$330,\"{m_code}\")")
        ws1.cell(idx, 5).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(idx, 5).font = Font(name=font_family, size=10)

        # Core Check
        c_core_cell = ws1.cell(idx, 6, m_core)
        c_core_cell.alignment = Alignment(horizontal="center", vertical="center")
        c_core_cell.font = Font(name=font_family, size=10, bold=(m_core == "✓"), color=c_core_txt if m_core == "✓" else "7F7F7F")
        if m_core == "✓":
            c_core_cell.fill = PatternFill(start_color=c_check_bg, end_color=c_check_bg, fill_type="solid")

        # Full Check
        c_full_cell = ws1.cell(idx, 7, m_full)
        c_full_cell.alignment = Alignment(horizontal="center", vertical="center")
        c_full_cell.font = Font(name=font_family, size=10, bold=True, color=c_green_txt)
        c_full_cell.fill = PatternFill(start_color=c_green_bg, end_color=c_green_bg, fill_type="solid")

        # Price
        p_cell = ws1.cell(idx, 8, m_price)
        p_cell.alignment = Alignment(horizontal="right", vertical="center")
        p_cell.font = Font(name=font_family, size=10)
        p_cell.number_format = '#,##0\\ "₫"'
        p_cell.fill = PatternFill(start_color=c_yellow_accent, end_color=c_yellow_accent, fill_type="solid")

        # Note
        n_cell = ws1.cell(idx, 9, m_note)
        n_cell.alignment = Alignment(horizontal="left", vertical="center")
        n_cell.font = Font(name=font_family, size=9, italic=True)

        for col_i in range(1, 10):
            ws1.cell(idx, col_i).border = cell_border

    # Set column widths for WS1
    widths_ws1 = {'A': 10.0, 'B': 16.0, 'C': 42.0, 'D': 46.0, 'E': 14.0, 'F': 12.0, 'G': 12.0, 'H': 20.0, 'I': 36.0}
    for col_letter, width in widths_ws1.items():
        ws1.column_dimensions[col_letter].width = width

    # =========================================================================
    # SHEET 2: 02_CHI_TIET_MODULE (Detailed Functional Requirements)
    # =========================================================================
    ws2 = wb.create_sheet(title="02_CHI_TIET_MODULE")
    ws2.views.sheetView[0].showGridLines = True

    # Title
    ws2.merge_cells("A1:G1")
    ws2["A1"] = "CHI TIẾT PHẠM VI CHỨC NĂNG HỌC VIỆN DUAXCAR KITCHEN – 01 MODULE = 01 QUY TRÌNH TRỌN GÓI"
    ws2["A1"].font = Font(name=font_family, size=14, bold=True, color=c_white)
    ws2["A1"].fill = PatternFill(start_color=c_navy, end_color=c_navy, fill_type="solid")
    ws2["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws2.row_dimensions[1].height = 36

    ws2.merge_cells("A2:G2")
    ws2["A2"] = "Bóc tách đầy đủ chi tiết từng đầu việc, luồng nghiệp vụ và tính năng của từng module trong hệ thống website DuaxCar Kitchen."
    ws2["A2"].font = Font(name=font_family, size=10, italic=True, color="595959")
    ws2["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws2.row_dimensions[2].height = 22

    # Headers (Row 4)
    headers_ws2 = ["Mã", "Module quy trình", "Loại", "Nhóm hạng mục", "Chi tiết chức năng / đầu việc", "Core", "Full"]
    for c_idx, h in enumerate(headers_ws2, 1):
        cell = ws2.cell(4, c_idx, h)
        cell.font = Font(name=font_family, size=10, bold=True, color=c_white)
        cell.fill = PatternFill(start_color=c_navy, end_color=c_navy, fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if c_idx in [1, 3, 6, 7] else "left", vertical="center")
        cell.border = cell_border
    ws2.row_dimensions[4].height = 26

    # Comprehensive Functional Line Items for DuaxCar Kitchen
    # Categorized into Trang/UI, UX/Responsive, Nghiệp vụ, Dữ liệu, Quản trị CMS, Tích hợp, Kỹ thuật, Bảo mật
    raw_features = [
        # --- M01: UX/UI Học viện Ẩm thực & Nền tảng Niềm tin (24 items) ---
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trang/UI", "Thiết kế Trang chủ theo nhận diện thương hiệu DuaxCar (Cam #FF7F00 & Đen #1A1A1A)", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trang/UI", "Thiết kế Header desktop + mobile, logo thương hiệu, menu điều hướng, hotline, CTA tư vấn", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trang/UI", "Thiết kế Footer: thông tin trung tâm, giấy phép kinh doanh, cơ sở bếp, mạng xã hội, chứng nhận", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trang/UI", "Thiết kế giao diện Danh mục khóa học tổng hợp & bộ lọc chuyên đề món ăn", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trang/UI", "Thiết kế giao diện Chi tiết khóa học đào tạo thực chiến & giáo trình chi tiết", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trang/UI", "Thiết kế giao diện Lịch khai giảng các lớp học sắp mở trong tháng", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trang/UI", "Thiết kế giao diện Đội ngũ Giảng viên & Master Chef ẩm thực Bàn tay vàng", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trang/UI", "Thiết kế giao diện Chi tiết hồ sơ Giảng viên (tiểu sử, giải thưởng, khóa dạy)", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trang/UI", "Thiết kế giao diện Cẩm nang ẩm thực & chia sẻ kinh nghiệm kinh doanh F&B", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trang/UI", "Thiết kế giao diện Chi tiết bài viết cẩm nang chuẩn SEO có Table of Contents tự động", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trang/UI", "Thiết kế Modal Form Đăng ký tư vấn khóa học nhanh & giữ chỗ ưu đãi", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trang/UI", "Thiết kế giao diện 404 Không tìm thấy trang & Trạng thái không có kết quả tìm kiếm", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "UX/Responsive", "Tối ưu hóa Mobile-first cho toàn bộ luồng tìm hiểu và đăng ký học trên điện thoại", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "UX/Responsive", "Responsive hoàn hảo trên các độ phân giải: Mobile, Tablet, Laptop, Màn hình lớn", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "UX/Responsive", "Nút CTA 'Đăng ký tư vấn ngay' và 'Xem lịch khai giảng' nổi bật, cố định trên mobile", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "UX/Responsive", "Hiển thị trạng thái Loading Skeleton, Submit form spinner, Thông báo kết quả tức thì", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "UX/Responsive", "Chuẩn hóa form nhập liệu, tự động validate số điện thoại Việt Nam 10 chữ số", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "UX/Responsive", "Thanh điều hướng Breadcrumb thông minh giúp học viên không bị mất phương hướng", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "UX/Responsive", "Nút liên hệ nhanh đa kênh cố định góc màn hình: Gọi Hotline, Nhắn Zalo OA, Fanpage", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "UX/Responsive", "Tích hợp tính năng chuyển đổi giao diện Sáng / Tối (Dark / Light Theme Toggle)", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trust UI", "Khối thông tin pháp nhân trung tâm đào tạo, mã số doanh nghiệp và cơ sở thực nghiệm", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trust UI", "Trust badges: Cam kết chuẩn vị truyền thống, Cầm tay chỉ việc 100%, Hỗ trợ mở quán", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trust UI", "Khối chứng nhận nghệ nhân ẩm thực, cúp vàng danh hiệu của đội ngũ giảng viên", "✓", "✓"),
        ("M01", "UX/UI Học viện Ẩm thực & Nền tảng Niềm tin", "CORE", "Trust UI", "Hiển thị liên kết chính sách đào tạo, quy định học phí và bảo lưu tại mọi form đăng ký", "✓", "✓"),

        # --- M02: CMS + Nội dung Thương hiệu, Quy chế & Pháp lý (24 items) ---
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "CMS quản trị", "Hệ thống quản lý trang nội dung tĩnh tập trung trong back-office", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "CMS quản trị", "Quản lý nội dung trang Giới thiệu Về DuaxCar: Lịch sử hình thành, Sứ mệnh, Tầm nhìn", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "CMS quản trị", "Quản lý hình ảnh và bài giới thiệu cơ sở vật chất khu vực bếp thực hành", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "CMS quản trị", "Quản lý bộ câu hỏi thường gặp (FAQ) phân theo nhóm chuyên đề tuyển sinh", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "CMS quản trị", "Thêm, sửa, sắp xếp thứ tự ưu tiên và xóa câu hỏi FAQ trực tiếp trên giao diện", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "CMS quản trị", "Quản lý Banner Slider / Hero Carousel trang chủ (ảnh, tiêu đề, mô tả, nút bấm, link)", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "CMS quản trị", "Quản lý nội dung thanh thông báo nổi bật (Top bar announcement) cho chiến dịch ưu đãi", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "CMS quản trị", "Quản lý thông tin liên hệ: Hotline chính, Hotline tư vấn 24/7, Email tuyển sinh, Giờ làm việc", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "CMS quản trị", "Quản lý địa chỉ các cơ sở đào tạo bếp và mã nhúng Google Maps định vị", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "CMS quản trị", "Quản lý liên kết mạng xã hội chính thức: Fanpage Facebook, TikTok, YouTube, Zalo", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "CMS quản trị", "Biên tập trực tiếp nội dung các trang chính sách pháp lý mà không cần can thiệp mã nguồn", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "CMS quản trị", "Cấu hình SEO Meta Title, Meta Description, Open Graph cho từng trang nội dung", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "Nội dung pháp lý", "Trang Điều khoản dịch vụ & Quy chế đào tạo của học viện ẩm thực", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "Nội dung pháp lý", "Trang Chính sách bảo mật thông tin cá nhân học viên theo quy định pháp luật", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "Nội dung pháp lý", "Trang Chính sách thanh toán học phí, quy định đặt cọc giữ chỗ và quy chế bảo lưu", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "Nội dung pháp lý", "Trang Thông tin pháp nhân doanh nghiệp, giấy phép kinh doanh và địa chỉ đăng ký", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "Nội dung thương hiệu", "Trang Về chúng tôi (Giới thiệu câu chuyện ẩm thực truyền thống & triết lý thực chiến)", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "Nội dung thương hiệu", "Trang Liên hệ & Đặt lịch tham quan trực tiếp xưởng bếp thực hành", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "Nội dung thương hiệu", "Trang Câu hỏi thường gặp (FAQ) tích hợp Accordion đóng/mở mượt mà", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "Nội dung thương hiệu", "Khối triết lý đào tạo: 'Học nghề thực chiến - Nắm công thức chuẩn - Tự tin mở quán'", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "Nội dung thương hiệu", "Khối cam kết quyền lợi học viên: Hỗ trợ tư vấn menu, setup bếp và bảo lưu học tập", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "Nội dung thương hiệu", "Khối chứng thực xã hội (Social Proof): Ý kiến đánh giá của học viên đã thành nghề", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "Nội dung thương hiệu", "Hệ thống Layout trang chính sách pháp lý chuẩn mực (PolicyLayout component)", "✓", "✓"),
        ("M02", "CMS + Nội dung Thương hiệu, Quy chế & Pháp lý", "CORE", "Nội dung thương hiệu", "Cơ chế Fallback nội dung mặc định đảm bảo trang web không bao giờ bị trống dữ liệu", "✓", "✓"),

        # --- M03: Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí (23 items) ---
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Catalogue/Admin", "Quản lý danh mục chuyên đề ẩm thực: Món ăn sáng, Lẩu nướng, Món ốc, Bếp trưởng...", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Catalogue/Admin", "Phân loại hình thức học: Trực tiếp tại xưởng bếp (Onsite) & Học Online/E-learning", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Catalogue/Admin", "Phân loại mục tiêu học: Học nghề mở quán kinh doanh trọn gói / Học 1-kèm-1 theo yêu cầu", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Catalogue/Admin", "Quản lý thông tin khóa học: Tên khóa học, Mã khóa, Slug SEO URL duy nhất", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Catalogue/Admin", "Cấu hình học phí niêm yết, giá ưu đãi hoặc tùy chọn 'Liên hệ nhận báo giá ưu đãi'", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Catalogue/Admin", "Cấu hình thời lượng khóa học: Số buổi học, tổng số giờ thực hành thực tế tại bếp", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Catalogue/Admin", "Giới hạn sĩ số học viên tối đa mỗi lớp (đảm bảo kèm cặp từng học viên)", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Catalogue/Admin", "Gán giảng viên / Master Chef phụ trách chính cho từng khóa học", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Catalogue/Admin", "Quản lý ảnh đại diện chất lượng cao cho từng khóa học", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Catalogue/Admin", "Quản lý bộ thẻ điểm nổi bật (Course Highlights): Công thức độc quyền, Hỗ trợ mở quán...", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Catalogue/Admin", "Bật/Tắt trạng thái Khóa học nổi bật (Featured Course) hiển thị trang chủ", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Catalogue/Admin", "Bật/Tắt trạng thái tuyển sinh của khóa học (Đang mở lớp / Tạm dừng tuyển sinh)", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "UI Khóa học", "Hiển thị Thẻ khóa học trực quan (Course Card) với đầy đủ thông tin tóm tắt", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "UI Khóa học", "Huy hiệu nhận diện hình thức học (Học tại bếp / E-learning / 1-kèm-1)", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "UI Khóa học", "Hiển thị mức học phí minh bạch hoặc nút gọi tư vấn học phí ưu đãi", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "UI Khóa học", "Hiển thị thông tin giảng viên đứng lớp và avatar Master Chef ngay trên thẻ khóa học", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Tìm kiếm & Lọc", "Tìm kiếm khóa học theo tên món, tên khóa học hoặc kỹ năng ẩm thực", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Tìm kiếm & Lọc", "Bộ lọc nhanh theo danh mục chuyên đề món ăn (Món sáng, Ốc & hải sản, Lẩu nướng...)", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Tìm kiếm & Lọc", "Bộ lọc theo hình thức đào tạo (Học tại bếp, Học mở quán, Học trực tuyến)", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Tìm kiếm & Lọc", "Sắp xếp khóa học theo: Khóa học nổi bật, Mới nhất, Học phí tăng/giảm dần", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Tìm kiếm & Lọc", "Trang hiển thị kết quả tìm kiếm với số lượng khóa học tìm thấy", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Tìm kiếm & Lọc", "Giao diện thân thiện khi không tìm thấy khóa học kèm gợi ý đăng ký tư vấn theo yêu cầu", "✓", "✓"),
        ("M03", "Catalogue Khóa học & Bộ lọc Tìm kiếm Đa tiêu chí", "CORE", "Tìm kiếm & Lọc", "Đồng bộ bộ lọc URL parameters giúp học viên dễ dàng chia sẻ liên kết tìm kiếm", "✓", "✓"),

        # --- M04: Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến (26 items) ---
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Trang chi tiết", "Trang chi tiết khóa học chuẩn cấu trúc tối ưu SEO (/khoa-hoc/[slug])", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Trang chi tiết", "Khối Hero khóa học: Tên khóa, chuyên mục, học phí ưu đãi, đánh giá sao từ học viên", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Trang chi tiết", "Thư viện hình ảnh thực tế các món ăn thành phẩm học viên sẽ được trực tiếp thực hành", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Trang chi tiết", "Video trailer / Giới thiệu khóa học và kỹ năng chế biến của Master Chef", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Trang chi tiết", "Bảng thông số lớp học: Thời lượng (số buổi), Giới hạn học viên, Hình thức thực hành", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Trang chi tiết", "Khối Điểm nổi bật (Highlights): Cung cấp công thức chuẩn gram, hướng dẫn tính giá cost", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Trang chi tiết", "Khối Quyền lợi học viên: 100% nguyên liệu tươi sạch, đồng phục, dụng cụ bếp tại chỗ", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Trang chi tiết", "Khối Cam kết đầu ra: Học lại miễn phí nếu chưa vững tay nghề, hỗ trợ sau khóa học", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Curriculum Accordion", "Giáo trình chi tiết tương tác dạng Accordion chia theo từng module / buổi học", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Curriculum Accordion", "Buổi 1: Lựa chọn, thẩm định và kỹ thuật sơ chế nguyên liệu tươi sống chuẩn nhà hàng", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Curriculum Accordion", "Buổi 2: Kỹ thuật xử lý nhiệt, hầm nước dùng trong ngọt và công thức pha chế sốt cốt độc quyền", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Curriculum Accordion", "Buổi 3: Thực hành chế biến chuyên sâu, kỹ thuật canh lửa, nêm nếm và ra món chuẩn tốc độ quán", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Curriculum Accordion", "Buổi 4: Tính toán giá cost nguyên liệu, định giá bán menu và thiết lập quy trình vận hành bếp quán", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Curriculum Accordion", "Hiển thị danh sách chi tiết các bài học, kỹ năng và công thức đạt được trong từng buổi", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Thông tin giảng viên", "Khối thông tin Giảng viên phụ trách lớp: Chân dung, danh hiệu, số năm kinh nghiệm", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Thông tin giảng viên", "Liên kết xem hồ sơ chi tiết và các giải thưởng ẩm thực của Master Chef", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Đăng ký khóa học", "Thanh CTA nổi bật cố định (Sticky Bar) hiển thị học phí và nút Đăng ký tư vấn", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Đăng ký khóa học", "Modal đăng ký tư vấn trực tiếp gắn liền với thông tin khóa học hiện tại", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Đăng ký khóa học", "Cơ chế tự động điền sẵn tên khóa học vào form đăng ký giúp học viên không phải nhập lại", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Khóa học liên quan", "Khối Khóa học đề xuất liên quan (Related Courses) hỗ trợ học viên chọn học theo combo", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Khóa học liên quan", "Gợi ý lộ trình đào tạo từ cơ bản đến nâng cao cho người khởi sự kinh doanh quán ăn", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Chia sẻ & Tương tác", "Nút chia sẻ thông tin khóa học nhanh lên Facebook, Zalo, sao chép liên kết", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "SEO Khóa học", "Cấu trúc dữ liệu có cấu trúc Course Schema theo tiêu chuẩn Schema.org của Google", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "SEO Khóa học", "Thẻ OpenGraph chuẩn hiển thị ảnh đại diện và học phí khi gửi link qua tin nhắn Zalo", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Curriculum Editor", "Trình chỉnh sửa giáo trình Curriculum JSON chuyên dụng trong CMS quản trị", "✓", "✓"),
        ("M04", "Chi tiết Khóa học & Giáo trình Đào tạo Thực chiến", "CORE", "Curriculum Editor", "Thêm, xóa, thay đổi thứ tự các buổi học linh hoạt mà không sợ sai lệch định dạng dữ liệu", "✓", "✓"),

        # --- M05: Đội ngũ Giảng viên & Master Chef Ẩm thực (16 items) ---
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "Hồ sơ chuyên gia", "Trang danh sách Đội ngũ Giảng viên & Nghệ nhân ẩm thực hàng đầu DuaxCar", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "Hồ sơ chuyên gia", "Trang chi tiết từng Giảng viên: Thầy Phạm Văn Long, Thầy Nguyễn Hữu Thọ, Thầy Lưu Đức Toàn", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "Hồ sơ chuyên gia", "Hiển thị danh hiệu cao quý: Nghệ nhân ẩm thực 'Bàn tay vàng', Bếp trưởng 25+ năm kinh nghiệm", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "Hồ sơ chuyên gia", "Ảnh chân dung chuyên nghiệp chất lượng cao của Master Chef", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "Hồ sơ chuyên gia", "Tiểu sử tóm tắt (Short Bio) và Tiểu sử chi tiết (Full Bio) về sự nghiệp ẩm thực", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "Hồ sơ chuyên gia", "Danh sách bằng cấp, chứng chỉ nghề bếp quốc gia và giải thưởng ẩm thực uy tín", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "Hồ sơ chuyên gia", "Triết lý giảng dạy & Câu trích dẫn tâm đắc (Quote) của giảng viên truyền cảm hứng", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "Hồ sơ chuyên gia", "Danh sách các khóa học do giảng viên trực tiếp đứng lớp và truyền nghề", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "CMS Giảng viên", "Hệ thống quản lý Giảng viên trong CMS back-office (/admin/giang-vien)", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "CMS Giảng viên", "Tạo mới, chỉnh sửa thông tin cá nhân, danh hiệu, chức vụ và số năm kinh nghiệm", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "CMS Giảng viên", "Quản lý danh sách thành tựu (Achievements JSON) dạng danh sách động", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "CMS Giảng viên", "Gán các khóa học phụ trách từ danh mục khóa học của hệ thống", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "CMS Giảng viên", "Tích hợp Media Selector để tải lên và chọn ảnh đại diện của giảng viên", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "SEO Giảng viên", "Tối ưu hóa SEO trang hồ sơ giảng viên (Person Schema chuẩn Google)", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "SEO Giảng viên", "Tăng cường chỉ số E-E-A-T (Chuyên môn, Thẩm quyền, Độ tin cậy) cho toàn bộ website", "✓", "✓"),
        ("M05", "Đội ngũ Giảng viên & Master Chef Ẩm thực", "CORE", "SEO Giảng viên", "Cơ chế bảo vệ thông tin cá nhân và bản quyền hình ảnh giảng viên", "✓", "✓"),

        # --- M06: Lịch Khai giảng & Quản lý Lớp học Tuyển sinh (16 items) ---
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "Lịch khai giảng", "Trang Lịch khai giảng công khai theo tháng trên website (/lich-khai-giang)", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "Lịch khai giảng", "Hiển thị danh sách các lớp học sắp mở trong tháng theo dòng thời gian trực quan", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "Lịch khai giảng", "Thông tin từng lớp: Tên khóa học, Ngày khai giảng chính xác, Địa điểm xưởng bếp", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "Lịch khai giảng", "Phân loại lịch học trong tuần: Thứ 2-4-6, Thứ 3-5-7 hoặc Lớp cấp tốc Thứ 7 & Chủ nhật", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "Lịch khai giảng", "Khung giờ ca học chi tiết: Ca sáng (8h30-11h30), Ca chiều (14h-17h), Ca tối (18h-21h)", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "Lịch khai giảng", "Trạng thái tuyển sinh trực quan: Đang mở nhận học viên / Sắp đủ chỗ (còn 1-2 chỗ) / Đã đóng", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "Lịch khai giảng", "Nút 'Đăng ký giữ chỗ ngay' gắn trực tiếp với lịch học và ca học được chọn", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "Lịch khai giảng", "Bộ lọc lịch khai giảng theo Chuyên mục món ăn hoặc theo Cơ sở đào tạo", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "Quản trị lịch học", "Quản lý danh sách lớp học khai giảng trong CMS quản trị", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "Quản trị lịch học", "Thiết lập ngày khai giảng, ca học, phòng bếp thực hành và giảng viên đứng lớp", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "Quản trị lịch học", "Cấu hình sĩ số học viên tối đa cho từng lớp để kiểm soát chất lượng đào tạo", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "Quản trị lịch học", "Tự động hoặc thủ công cập nhật trạng thái lớp khi số học viên đăng ký đạt giới hạn", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "Quản trị lịch học", "Thông báo cảnh báo khi lớp học chuẩn bị đến ngày khai giảng", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "Quản trị lịch học", "Lưu trữ lịch sử các lớp đã hoàn thành phục vụ công tác tra cứu đào tạo", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "UX/Hiệu ứng", "Giao diện thẻ lịch khai giảng nổi bật, dễ nhìn trên thiết bị di động", "✓", "✓"),
        ("M06", "Lịch Khai giảng & Quản lý Lớp học Tuyển sinh", "CORE", "UX/Hiệu ứng", "Cơ chế nhắc nhở học viên đăng ký sớm để nhận ưu đãi học phí", "✓", "✓"),

        # --- M07: Thu thập Lead & Quy trình Tiếp nhận Đăng ký (20 items) ---
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Form đăng ký", "Form Đăng ký tư vấn khóa học nhanh tại Trang chủ và chân trang các bài viết", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Form đăng ký", "Modal Đăng ký tư vấn chuyên nghiệp mở từ mọi nút CTA trên toàn hệ thống", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Form đăng ký", "Form Liên hệ & Đặt lịch tư vấn tại trang Liên hệ (/lien-he)", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Trường thông tin", "Thu thập Họ và tên học viên", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Trường thông tin", "Thu thập Số điện thoại (tự động kiểm tra cú pháp 10 số, đầu số hợp lệ của nhà mạng VN)", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Trường thông tin", "Thu thập Email học viên để gửi xác nhận và tài liệu", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Trường thông tin", "Lựa chọn Khóa học quan tâm từ danh sách khóa học đang hoạt động", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Trường thông tin", "Lựa chọn Cơ sở bếp thực hành mong muốn tham gia học", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Trường thông tin", "Lựa chọn Nhu cầu cụ thể: Mở quán kinh doanh / Nâng cao tay nghề / Nấu ăn gia đình", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Trường thông tin", "Ô nhập ghi chú nhu cầu cá nhân hoặc câu hỏi muốn chuyên gia giải đáp", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Xử lý dữ liệu", "Validate dữ liệu chặt chẽ ở cả phía Client (React Hook Form/Zod) và Server API", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Xử lý dữ liệu", "Cơ chế chống gửi trùng lặp (Submit Debounce) khi học viên bấm nút nhiều lần", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Xử lý dữ liệu", "Cơ chế phòng chống spam tự động (Honeypot field & Rate limiting API)", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Xử lý dữ liệu", "Lưu trữ thông tin đăng ký an toàn vào bảng registrations trong Supabase PostgreSQL", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Xử lý dữ liệu", "Ghi nhận nguồn đăng ký (URL trang đăng ký, thời gian tạo, thiết bị)", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Phản hồi học viên", "Màn hình popup thông báo Đăng ký thành công tức thì với giao diện chuyên nghiệp", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Phản hồi học viên", "Thông báo rõ ràng quy trình tiếp theo: Chuyên viên tuyển sinh sẽ gọi lại trong 15-30 phút", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "Phản hồi học viên", "Cung cấp Hotline hỗ trợ khẩn cấp nếu học viên cần giải đáp ngay lập tức", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "API Tuyển sinh", "Route API /api/contact tiếp nhận thông tin liên hệ và form tư vấn", "✓", "✓"),
        ("M07", "Thu thập Lead & Quy trình Tiếp nhận Đăng ký", "CORE", "API Tuyển sinh", "Route API /api/cms/registrations xử lý dữ liệu đăng ký khóa học", "✓", "✓"),

        # --- M08: Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM (22 items) ---
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Admin Dashboard", "Giao diện quản trị CMS tổng quan (/admin) hiện đại, trực quan, hỗ trợ Dark/Light mode", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Admin Dashboard", "Thẻ thống kê tổng số đơn đăng ký học viên mới tiếp nhận trong ngày/tuần/tháng", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Admin Dashboard", "Thống kê tổng số khóa học đang mở tuyển sinh và số bài viết cẩm nang", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Admin Dashboard", "Bảng danh sách các học viên mới đăng ký cần liên hệ tư vấn gấp (chờ xử lý)", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Lead CRM", "Trang quản lý đơn đăng ký tuyển sinh chuyên dụng (/admin/dang-ky)", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Lead CRM", "Bảng hiển thị danh sách lead với đầy đủ: Họ tên, Số điện thoại, Email, Khóa học, Ngày gửi", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Lead CRM", "Bộ lọc đơn đăng ký theo Trạng thái xử lý: Mới tiếp nhận, Đã liên hệ, Đã nhập học, Hủy", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Lead CRM", "Bộ lọc đơn đăng ký theo từng Khóa học cụ thể phục vụ giáo vụ tổng hợp danh sách lớp", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Lead CRM", "Tìm kiếm nhanh học viên theo Số điện thoại hoặc Họ tên", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Lead CRM", "Cập nhật Trạng thái xử lý đơn đăng ký (Pending -> Contacted -> Enrolled -> Cancelled)", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Lead CRM", "Ghi chú chăm sóc khách hàng nội bộ (lưu lại nhu cầu đặc biệt, hẹn ngày nhập học)", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Lead CRM", "Lưu vết thời gian cập nhật trạng thái đơn phục vụ quản lý hiệu suất tư vấn", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Quản lý Khóa học", "Trang quản lý khóa học (/admin/khoa-hoc): Tạo mới, sửa, nhân bản và xóa khóa học", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Quản lý Khóa học", "Form chỉnh sửa thông tin khóa học: Tên, slug, danh mục, hình thức học, học phí, sĩ số", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Quản lý Khóa học", "Trình chỉnh sửa giáo trình chi tiết (Curriculum Builder) theo từng buổi học", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Quản lý Khóa học", "Quản lý thẻ điểm nổi bật (Highlights tags) và gán Giảng viên phụ trách", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Quản lý Giảng viên", "Trang quản lý giảng viên (/admin/giang-vien): Tạo mới, cập nhật hồ sơ Master Chef", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Quản lý Giảng viên", "Cập nhật danh hiệu, số năm kinh nghiệm, giải thưởng, thành tựu và câu nói tâm đắc", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Cài đặt & Pháp lý", "Trang quản lý FAQ (/admin/faq) thêm sửa xóa câu hỏi thường gặp", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Cài đặt & Pháp lý", "Trang quản lý Chính sách (/admin/chinh-sach) cập nhật quy chế đào tạo", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Cài đặt & Pháp lý", "Trang cài đặt hệ thống (/admin/cai-dat) cấu hình thương hiệu, hotline, mạng xã hội", "✓", "✓"),
        ("M08", "Back-office Quản trị Tuyển sinh, Khóa học & Lead CRM", "CORE", "Bảo mật quản trị", "Xác thực đăng nhập quản trị qua Supabase Auth, bảo vệ tuyến đường /admin/*", "✓", "✓"),

        # --- M09: Kho Media Tập trung & Tối ưu Nén ảnh Client-side (16 items) ---
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Quản lý Media", "Trang thư viện Media tập trung (/admin/media) quản lý toàn bộ hình ảnh và video", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Nén ảnh tự động", "Tích hợp thư viện nén ảnh tự động phía Client trước khi upload (image-compressor.ts)", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Nén ảnh tự động", "Tự động chuyển đổi hình ảnh sang định dạng WebP hiện đại, giảm dung lượng 70-80%", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Nén ảnh tự động", "Giữ nguyên độ sắc nét và màu sắc trung thực của món ăn phục vụ hiển thị chất lượng cao", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Tải lên & Lưu trữ", "Hỗ trợ tải lên nhiều hình ảnh cùng lúc (Batch Upload) với thanh tiến trình trực quan", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Tải lên & Lưu trữ", "Lưu trữ tài nguyên an toàn trên Supabase Storage hoặc Cloud Storage tối ưu băng thông", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Tải lên & Lưu trữ", "Route API /api/cms/upload xử lý xác thực và kiểm tra dung lượng, định dạng file", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Bộ chọn ảnh Modal", "Modal chọn ảnh dùng chung (Media Picker Modal) tái sử dụng trong toàn bộ CMS", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Bộ chọn ảnh Modal", "Tích hợp Media Picker vào form Khóa học (chọn ảnh bìa và gallery món ăn)", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Bộ chọn ảnh Modal", "Tích hợp Media Picker vào form Giảng viên (chọn ảnh chân dung Master Chef)", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Bộ chọn ảnh Modal", "Tích hợp Media Picker vào trình soạn thảo bài viết Blog / Tin tức", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Bộ chọn ảnh Modal", "Tích hợp Media Picker vào Cài đặt hệ thống (chọn Logo, Favicon, Banner Hero)", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Xem trước & Tiện ích", "Xem trước ảnh kích thước lớn bằng Lightbox Modal không cần tải file về máy", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Xem trước & Tiện ích", "Tính năng Copy nhanh đường dẫn URL ảnh vào clipboard với 1 cú click", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Xem trước & Tiện ích", "Xóa hình ảnh không còn sử dụng kèm hộp thoại xác nhận an toàn", "✓", "✓"),
        ("M09", "Kho Media Tập trung & Tối ưu Nén ảnh Client-side", "CORE", "Xem trước & Tiện ích", "Phân trang và tìm kiếm hình ảnh theo tên file trong thư viện media", "✓", "✓"),

        # --- M10: Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B (22 items) ---
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang danh mục", "Trang danh sách cẩm nang ẩm thực & tin tức chuyên ngành F&B (/tin-tuc)", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang danh mục", "Phân loại bài viết theo chuyên mục: Bí quyết nấu ăn ngon, Kinh nghiệm mở quán, Xu hướng F&B", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang danh mục", "Hiển thị bài viết tiêu biểu (Featured Post) với kích thước lớn nổi bật ở đầu trang", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang danh mục", "Thẻ bài viết hiển thị ảnh bìa, tiêu đề, đoạn trích ngắn, tác giả, ngày đăng và thời gian đọc", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang danh mục", "Tìm kiếm bài viết cẩm nang theo từ khóa và lọc theo chuyên mục", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang chi tiết", "Trang chi tiết bài viết cẩm nang chuẩn SEO (/tin-tuc/[slug])", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang chi tiết", "Mục lục tự động (Table of Contents - TOC) tự động nhận diện thẻ H2, H3 trong bài viết", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang chi tiết", "TOC tự động scroll đến phần nội dung tương ứng và highlight mục đang đọc", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang chi tiết", "Hiển thị thông tin tác giả bài viết (chuyên gia Master Chef hoặc ban biên tập DuaxCar)", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang chi tiết", "Hiển thị thời gian đọc ước tính (ví dụ: '5 phút đọc') giúp nâng cao trải nghiệm người dùng", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang chi tiết", "Trình bày nội dung đẹp mắt với kiểu chữ chuyên nghiệp, khoảng cách dòng thoáng đãng", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang chi tiết", "Chèn ảnh minh họa món ăn, bảng định lượng nguyên liệu và video hướng dẫn vào bài viết", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang chi tiết", "Khối bài viết liên quan (Related Articles) giúp giữ chân độc giả ở lại website", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang chi tiết", "Nút chia sẻ bài viết lên Facebook, Zalo, Twitter và sao chép link", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Trang chi tiết", "Khối CTA giới thiệu khóa học ẩm thực liên quan đặt ở cuối bài viết", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Quản trị bài viết", "Trang quản trị bài viết trong CMS (/admin/tin-tuc): Tạo mới, sửa, xóa bài viết", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Quản trị bài viết", "Trình soạn thảo văn bản phong phú (Rich Text Editor) trực quan, dễ dùng", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Quản trị bài viết", "Định dạng tiêu đề H2, H3, bôi đậm, in nghiêng, gạch đầu dòng, căn lề, trích dẫn", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Quản trị bài viết", "Chèn hình ảnh trực tiếp từ thư viện Media vào vị trí con trỏ chuột trong bài", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "Quản trị bài viết", "Tùy biến Slug URL thân thiện, trích đoạn ngắn Excerpt và cấu hình bài viết nổi bật", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "SEO Bài viết", "Cấu trúc dữ liệu có cấu trúc Article Schema theo tiêu chuẩn Google cho bài viết", "✓", "✓"),
        ("M10", "Blog Cẩm nang Ẩm thực, Bí quyết Nấu ăn & Mở quán F&B", "CORE", "SEO Bài viết", "Tự động tạo thẻ OpenGraph Title, Description, Image phục vụ chia sẻ bài viết mạng xã hội", "✓", "✓"),

        # --- M11: SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking (20 items - ADD-ON FULL) ---
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "SEO On-page", "Tự động sinh Meta Tags đầy đủ cho toàn bộ trang: Title, Description, Canonical URL", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "SEO On-page", "Tối ưu hóa tiêu đề và mô tả theo bộ từ khóa ẩm thực truyền thống & học nghề mở quán", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "SEO Schema", "Cấu hình Course Schema cho tất cả các trang khóa học (tên, học phí, giảng viên, đơn vị)", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "SEO Schema", "Cấu hình EducationalOrganization Schema định danh DuaxCar là đơn vị đào tạo uy tín", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "SEO Schema", "Cấu hình Person Schema cho hồ sơ Master Chef và đội ngũ nghệ nhân ẩm thực", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "SEO Schema", "Cấu hình Article Schema và FAQPage Schema cho trang tin tức và hỏi đáp", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "SEO Schema", "Cấu hình BreadcrumbList Schema giúp Google hiển thị cây điều hướng trên kết quả tìm kiếm", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "SEO Technical", "Tự động tạo sitemap.xml động cập nhật theo thời gian thực khi có khóa học/bài viết mới", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "SEO Technical", "Cấu hình chuẩn file robots.txt chỉ dẫn công cụ tìm kiếm thu thập dữ liệu hiệu quả", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "SEO Technical", "Cấu hình thẻ Canonical tránh trùng lặp nội dung giữa các phiên bản URL", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "Social Share", "Tối ưu Open Graph tags & Twitter Cards hiển thị thumbnail đẹp mắt khi share link", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "Analytics", "Tích hợp Google Analytics 4 (GA4) theo dõi chi tiết lưu lượng truy cập và hành vi học viên", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "Analytics", "Tích hợp Google Tag Manager (GTM) quản lý tập trung toàn bộ mã theo dõi tiếp thị", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "Tracking chuyển đổi", "Cài đặt theo dõi sự kiện chuyển đổi Xem Khóa Học (view_course_detail)", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "Tracking chuyển đổi", "Cài đặt theo dõi sự kiện Bắt đầu điền form đăng ký (begin_course_registration)", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "Tracking chuyển đổi", "Cài đặt theo dõi sự kiện Đăng ký khóa học thành công (complete_registration)", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "Tracking chuyển đổi", "Cài đặt theo dõi sự kiện Nhấp gọi Hotline và Nhấp chat Zalo OA (click_to_contact)", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "Advertising Pixel", "Tích hợp Meta Pixel (Facebook Pixel) tối ưu chiến dịch chạy quảng cáo tuyển sinh học viên", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "Advertising Pixel", "Cấu hình Custom Conversion trên Facebook Events Manager đo lường chi phí trên mỗi lead", "—", "✓"),
        ("M11", "SEO Chuyên sâu Ngành Đào tạo Ẩm thực & Analytics Tracking", "ADD-ON FULL", "Advertising Pixel", "Tích hợp TikTok Pixel (nếu trung tâm chạy quảng cáo video nấu ăn trên TikTok)", "—", "✓"),

        # --- M12: Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành Tối ưu (20 items - ADD-ON FULL) ---
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Hạ tầng Cloud", "Triển khai hệ thống trên nền tảng Vercel Edge Serverless / Cloud Hosting tốc độ cao", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Hạ tầng Cloud", "Cấu hình môi trường Production và Staging riêng biệt phục vụ thử nghiệm tính năng", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Cơ sở dữ liệu", "Cơ sở dữ liệu Supabase Managed PostgreSQL hiệu năng cao, chịu tải lớn", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Cơ sở dữ liệu", "Cơ chế Ping Keep-Alive tự động (/api/keep-alive) giữ kết nối database 24/7 không bị idle", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Cơ sở dữ liệu", "Tự động sao lưu dữ liệu (Daily Database Backup) định kỳ bảo đảm an toàn dữ liệu", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Bảo mật dữ liệu", "Kích hoạt hệ thống Row Level Security (RLS) trên toàn bộ các bảng trong database", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Bảo mật dữ liệu", "Bảo vệ thông tin cá nhân của học viên, ngăn chặn truy cập trái phép từ bên ngoài", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Bảo mật xác thực", "Hệ thống xác thực quản trị viên qua Supabase Auth với session cookie mã hóa an toàn", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Bảo mật ứng dụng", "Chống tấn công CSRF, XSS và lọc dữ liệu SQL Injection tự động bằng ORM", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Bảo mật ứng dụng", "Ẩn biến môi trường bảo mật (.env.production) chứa API keys và thông tin kết nối DB", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Tối ưu hiệu năng", "Áp dụng Next.js Turbopack & React Server Components tối ưu tốc độ tải trang dưới 1.5s", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Tối ưu hiệu năng", "Áp dụng Incremental Static Regeneration (ISR) tự động cache trang khóa học và blog", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Tối ưu hiệu năng", "Tối ưu hóa hình ảnh tự động qua Next Image (WebP/AVIF, lazy loading, responsive size)", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Tối ưu hiệu năng", "Đạt điểm Google Lighthouse > 90 trên cả phiên bản Desktop và Mobile", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "CDN & SSL", "Tích hợp Cloudflare CDN phân phối nội dung tĩnh toàn cầu với độ trễ cực thấp", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "CDN & SSL", "Chứng chỉ bảo mật SSL/TLS miễn phí trọn đời (HTTPS xanh uy tín)", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Giám sát & Log", "Theo dõi trạng thái hệ thống và cảnh báo lỗi thời gian thực (Error Logging)", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Giám sát & Log", "Giám sát lưu lượng băng thông và dung lượng lưu trữ trên Cloud", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Vận hành", "Bàn giao tài liệu hướng dẫn quản trị và vận hành kỹ thuật chi tiết", "—", "✓"),
        ("M12", "Hạ tầng Cloud Cao cấp, Bảo mật Database & Vận hành", "ADD-ON FULL", "Vận hành", "Cam kết SLA hỗ trợ kỹ thuật và xử lý sự cố trong suốt thời gian bảo hành", "—", "✓"),

        # --- M13: Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến (18 items - ADD-ON FULL) ---
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Thanh toán VietQR", "Tự động sinh mã VietQR động theo tiêu chuẩn Napas cho từng đơn đăng ký khóa học", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Thanh toán VietQR", "Mã QR nhúng chính xác số tiền cọc (ví dụ 1.000.000đ - 2.000.000đ) hoặc toàn bộ học phí", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Thanh toán VietQR", "Nội dung chuyển khoản được tạo tự động với mã định danh học viên: DUAXCAR [MãĐơn] [SĐT]", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Thanh toán VietQR", "Học viên chỉ cần quét mã bằng app ngân hàng là hoàn tất, không sợ chuyển nhầm số tiền", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Cổng thanh toán", "Tích hợp cổng thanh toán trực tuyến (VNPay / MoMo / ZaloPay / OnePay)", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Cổng thanh toán", "Hỗ trợ thanh toán qua Thẻ ATM nội địa (Internet Banking của tất cả ngân hàng)", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Cổng thanh toán", "Hỗ trợ thanh toán qua Thẻ quốc tế Visa, Mastercard, JCB", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Cổng thanh toán", "Hỗ trợ thanh toán qua Ví điện tử MoMo, ZaloPay, Viettel Money", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Webhook & Đối soát", "Webhook tiếp nhận tín hiệu thanh toán tự động thời gian thực từ Cổng thanh toán/Ngân hàng", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Webhook & Đối soát", "Tự động cập nhật trạng thái đơn đăng ký sang 'Đã đặt cọc' hoặc 'Đã đóng đủ học phí'", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Webhook & Đối soát", "Tự động khóa slot và trừ số lượng chỗ còn lại của lớp học khi nhận cọc thành công", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Giao diện thanh toán", "Màn hình hướng dẫn thanh toán chi tiết kèm đồng hồ đếm ngược giữ chỗ (15 phút)", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Giao diện thanh toán", "Màn hình thông báo Xác nhận thanh toán cọc thành công kèm mã biên nhận điện tử", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Giao diện thanh toán", "Màn hình thông báo Thanh toán thất bại / hết hạn giữ chỗ kèm nút thử lại", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Quản lý thanh toán", "Module quản lý giao dịch học phí trong CMS quản trị", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Quản lý thanh toán", "Quản lý danh sách biên lai thu cọc, lọc theo ngày, theo khóa học, theo phương thức", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Quản lý thanh toán", "Hỗ trợ nhân viên kế toán xác nhận thanh toán thủ công đối với trường hợp chuyển khoản trực tiếp", "—", "✓"),
        ("M13", "Cổng Đặt cọc Giữ chỗ Khóa học & Cổng Học phí Trực tuyến", "ADD-ON FULL", "Quản lý thanh toán", "Xuất file danh sách giao dịch nộp học phí ra Excel phục vụ công tác kế toán", "—", "✓"),

        # --- M14: Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công (18 items - ADD-ON FULL) ---
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Showcase thành công", "Trang chuyên đề 'Gương mặt Khởi nghiệp F&B - Học viên mở quán thành công'", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Showcase thành công", "Showcase hình ảnh thực tế quán ăn, nhà hàng, quầy đồ ăn sáng do học viên làm chủ", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Showcase thành công", "Video phỏng vấn học viên: Quá trình học nghề, khó khăn khi setup và thành quả mở quán", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Showcase thành công", "Hiển thị doanh thu thực tế, số lượng tô/suất bán ra mỗi ngày để truyền cảm hứng", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Showcase thành công", "Địa chỉ quán ăn của học viên kèm link Google Maps (giúp học viên quảng bá quán mới)", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Đánh giá học viên", "Hệ thống Đánh giá & Cảm nhận học viên có xác thực (Verified Student Reviews)", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Đánh giá học viên", "Đánh giá xếp hạng sao (Rating 1-5 sao) theo từng khóa học cụ thể", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Đánh giá học viên", "Học viên được tải lên hình ảnh món ăn tự nấu tại nhà/quán chứng minh chất lượng", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Đánh giá học viên", "Hiển thị điểm đánh giá trung bình và tỷ lệ hài lòng (ví dụ: 4.9/5 sao từ 200+ học viên)", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Đánh giá học viên", "Huy hiệu 'Học viên đã tốt nghiệp' gắn liền với từng nhận xét", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Kiểm duyệt CMS", "Module quản lý và kiểm duyệt đánh giá học viên trong CMS", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Kiểm duyệt CMS", "Cơ chế kiểm duyệt trước khi hiển thị (Moderation Queue) chống spam nhận xét tiêu cực giả mạo", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Kiểm duyệt CMS", "Ban quản trị phản hồi trực tiếp dưới từng nhận xét của học viên", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Hiển thị thông minh", "Widget hiển thị đánh giá học viên tại trang chủ và nhúng trong trang chi tiết khóa học", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Hiển thị thông minh", "Bộ lọc đánh giá: Xem đánh giá có ảnh, lọc theo số sao (5 sao, 4 sao)", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Hiển thị thông minh", "SEO Review Schema giúp hiển thị ngôi sao đánh giá màu vàng trên kết quả tìm kiếm Google", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Chứng chỉ tốt nghiệp", "Khu vực trưng bày mẫu Chứng chỉ hoàn thành khóa đào tạo nghề ẩm thực DuaxCar", "—", "✓"),
        ("M14", "Social Proof Chuyên sâu & Showcase Học viên Mở quán Thành công", "ADD-ON FULL", "Chứng chỉ tốt nghiệp", "Hệ thống tra cứu mã chứng chỉ học viên điện tử chống làm giả", "—", "✓"),

        # --- M15: Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền (16 items - ADD-ON FULL) ---
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Cổng học viên", "Hệ thống Đăng nhập / Đăng ký tài khoản học viên bằng Số điện thoại / Email / OTP", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Cổng học viên", "Dashboard cá nhân học viên (/hoc-vien): Quản lý thông tin cá nhân và lớp học", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Cổng học viên", "Danh sách các khóa học đã đăng ký, lịch khai giảng và thông tin phòng bếp thực hành", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Cổng học viên", "Theo dõi tình trạng học phí: Đã cọc, Còn thiếu, Đã thanh toán trọn gói", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Kho công thức số", "Kho tài liệu công thức điện tử chuẩn định lượng (Digital Recipe Hub)", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Kho công thức số", "Phân quyền truy cập tài liệu: Chỉ học viên đã đóng đủ học phí khóa nào mới xem được khóa đó", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Kho công thức số", "File PDF công thức chuẩn gram/ml chi tiết từng nguyên liệu và tỷ lệ gia vị sốt cốt", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Kho công thức số", "Video clip quay cận cảnh kỹ thuật chế biến thao tác mẫu của Master Chef", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Kho công thức số", "Bảng tính Excel mẫu định lượng giá cost nguyên liệu và tính toán điểm hòa vốn mở quán", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Kho công thức số", "Danh bạ nhà cung cấp nguyên liệu, gia vị sỉ chuẩn vị được DuaxCar liên kết trợ giá", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Bảo vệ bản quyền", "Cơ chế đóng dấu Watermark tự động (Họ tên + SĐT học viên) trên toàn bộ tài liệu PDF", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Bảo vệ bản quyền", "Chống sao chép, tải lậu và phát tán bí quyết công thức độc quyền ra ngoài", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Hỗ trợ học viên", "Học viên gửi câu hỏi thắc mắc trực tiếp cho Giảng viên Master Chef qua cổng cá nhân", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Hỗ trợ học viên", "Gửi yêu cầu đăng ký học lại miễn phí hoặc xin bảo lưu lịch học khi bận việc đột xuất", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Quản lý Cổng", "Module quản lý tài khoản học viên và phân quyền khóa học trong CMS quản trị", "—", "✓"),
        ("M15", "Cổng Học viên (Student Portal) & Kho Công thức Số Độc quyền", "ADD-ON FULL", "Quản lý Cổng", "Khóa / Mở quyền truy cập tài liệu công thức theo tiến độ học của từng học viên", "—", "✓"),

        # --- M16: CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp (18 items - ADD-ON FULL) ---
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Email Automation", "Tích hợp dịch vụ gửi Transactional Email tự động (Resend / Brevo)", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Email Automation", "Tự động gửi Thư xác nhận đăng ký tư vấn kèm lộ trình khóa học ngay khi học viên điền form", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Email Automation", "Tự động gửi Thư xác nhận giữ chỗ thành công khi nhận được chuyển khoản cọc", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Email Automation", "Tự động gửi Email nhắc lịch khai giảng và hướng dẫn trang phục/dụng cụ trước ngày học 02 ngày", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Email Automation", "Thư chúc mừng tốt nghiệp và gửi link đánh giá chất lượng khóa học sau buổi bế giảng", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "ZNS / SMS", "Tích hợp tin nhắn Zalo ZNS / SMS Brandname gửi thông báo tuyển sinh", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "ZNS / SMS", "Gửi tin nhắn ZNS xác nhận mã giữ chỗ và số điện thoại hotline của chuyên viên phụ trách", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "ZNS / SMS", "Gửi tin nhắn SMS nhắc lịch khai giảng sáng sớm ngày đầu tiên vào bếp", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Webhook & Tích hợp", "Đồng bộ dữ liệu Lead tuyển sinh tự động sang Google Sheets thời gian thực phục vụ telesale", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Webhook & Tích hợp", "Webhook đồng bộ dữ liệu sang phần mềm CRM ngoài (HubSpot / Lark Suite / Bitrix24)", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Phân quyền quản trị", "Phân quyền người dùng quản trị đa cấp (Multi-role CMS) chuyên biệt cho từng bộ phận", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Phân quyền quản trị", "Nhóm Super Admin: Toàn quyền cấu hình hệ thống, khóa học, cài đặt và doanh thu", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Phân quyền quản trị", "Nhóm Chuyên viên Tuyển sinh: Chỉ xem và chăm sóc danh sách Lead học viên được giao", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Phân quyền quản trị", "Nhóm Kế toán: Chỉ xem, xác nhận và đối soát thanh toán cọc học phí", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Phân quyền quản trị", "Nhóm Biên tập viên Content: Chỉ tạo và chỉnh sửa bài viết cẩm nang ẩm thực", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Báo cáo tuyển sinh", "Báo cáo thống kê số lượng lead theo từng kênh (Website, Facebook Ads, Giới thiệu)", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Báo cáo tuyển sinh", "Báo cáo tỷ lệ chuyển đổi từ đăng ký sang nộp cọc và nhập học chính thức", "—", "✓"),
        ("M16", "CRM Tuyển sinh Tự động, Automation Thông báo (Email / ZNS) & Vận hành Đa cấp", "ADD-ON FULL", "Audit Log", "Nhật ký hoạt động (Audit Log) ghi nhận toàn bộ thao tác cập nhật dữ liệu của nhân viên", "—", "✓")
    ]

    total_core_items = 0
    total_full_items = 0

    for r_idx, item in enumerate(raw_features, 5):
        ws2.row_dimensions[r_idx].height = 20
        m_code, m_name, m_type, m_group, m_desc, m_core, m_full = item
        
        ws2.cell(r_idx, 1, m_code).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(r_idx, 1).font = Font(name=font_family, size=9, bold=True, color=c_core_txt)
        
        ws2.cell(r_idx, 2, m_name).alignment = Alignment(horizontal="left", vertical="center")
        ws2.cell(r_idx, 2).font = Font(name=font_family, size=9, bold=True, color=c_core_txt if m_type=="CORE" else c_green_txt)
        
        ws2.cell(r_idx, 3, m_type).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(r_idx, 3).font = Font(name=font_family, size=8, bold=True, color=c_core_txt if m_type=="CORE" else c_green_txt)
        
        ws2.cell(r_idx, 4, m_group).alignment = Alignment(horizontal="left", vertical="center")
        ws2.cell(r_idx, 4).font = Font(name=font_family, size=9)
        
        ws2.cell(r_idx, 5, m_desc).alignment = Alignment(horizontal="left", vertical="center")
        ws2.cell(r_idx, 5).font = Font(name=font_family, size=9)
        
        # Core check
        c_cell = ws2.cell(r_idx, 6, m_core)
        c_cell.alignment = Alignment(horizontal="center", vertical="center")
        c_cell.font = Font(name=font_family, size=10, bold=(m_core=="✓"), color=c_core_txt if m_core=="✓" else "7F7F7F")
        if m_core == "✓":
            total_core_items += 1
            c_cell.fill = PatternFill(start_color=c_check_bg, end_color=c_check_bg, fill_type="solid")
            
        # Full check
        f_cell = ws2.cell(r_idx, 7, m_full)
        f_cell.alignment = Alignment(horizontal="center", vertical="center")
        f_cell.font = Font(name=font_family, size=10, bold=True, color=c_green_txt)
        if m_full == "✓":
            total_full_items += 1
            f_cell.fill = PatternFill(start_color=c_green_bg, end_color=c_green_bg, fill_type="solid")

        # Alternating row fill
        if r_idx % 2 == 0:
            for col_i in range(1, 6):
                ws2.cell(r_idx, col_i).fill = PatternFill(start_color=c_alt_row, end_color=c_alt_row, fill_type="solid")

        for col_i in range(1, 8):
            ws2.cell(r_idx, col_i).border = cell_border

    last_feature_row = len(raw_features) + 4
    summary_row = last_feature_row + 1
    
    # Add Total row in Sheet 2
    ws2.cell(summary_row, 4, "TỔNG SỐ HẠNG MỤC TÍNH NĂNG").alignment = Alignment(horizontal="right", vertical="center")
    ws2.cell(summary_row, 4).font = Font(name=font_family, size=10, bold=True, color=c_navy)
    
    ws2.cell(summary_row, 6, f"=COUNTIF(F5:F{last_feature_row},\"✓\")").alignment = Alignment(horizontal="center", vertical="center")
    ws2.cell(summary_row, 6).font = Font(name=font_family, size=11, bold=True, color=c_core_txt)
    ws2.cell(summary_row, 6).fill = PatternFill(start_color=c_gray_bg, end_color=c_gray_bg, fill_type="solid")
    
    ws2.cell(summary_row, 7, f"=COUNTIF(G5:G{last_feature_row},\"✓\")").alignment = Alignment(horizontal="center", vertical="center")
    ws2.cell(summary_row, 7).font = Font(name=font_family, size=11, bold=True, color=c_green_txt)
    ws2.cell(summary_row, 7).fill = PatternFill(start_color=c_green_bg, end_color=c_green_bg, fill_type="solid")
    
    for col_i in range(1, 8):
        ws2.cell(summary_row, col_i).border = thick_bottom

    # Update Sheet 1 formula to point to the summary row
    ws1["E6"] = f"='02_CHI_TIET_MODULE'!F{summary_row}"
    ws1["H6"] = f"='02_CHI_TIET_MODULE'!G{summary_row}"

    # Set column widths for WS2
    widths_ws2 = {'A': 8.0, 'B': 42.0, 'C': 16.0, 'D': 24.0, 'E': 75.0, 'F': 11.0, 'G': 11.0}
    for col_letter, width in widths_ws2.items():
        ws2.column_dimensions[col_letter].width = width

    # =========================================================================
    # SHEET 3: 03_TRANG_UI (UI Screen Inventory)
    # =========================================================================
    ws3 = wb.create_sheet(title="03_TRANG_UI")
    ws3.views.sheetView[0].showGridLines = True

    # Title
    ws3.merge_cells("A1:F1")
    ws3["A1"] = "DANH SÁCH TRANG / MÀN HÌNH UI THỰC TẾ BÀN GIAO – DUAXCAR KITCHEN"
    ws3["A1"].font = Font(name=font_family, size=14, bold=True, color=c_white)
    ws3["A1"].fill = PatternFill(start_color=c_navy, end_color=c_navy, fill_type="solid")
    ws3["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws3.row_dimensions[1].height = 36

    ws3.merge_cells("A2:F2")
    ws3["A2"] = "Gói Core: 29 trang/màn hình. Gói Full: 40 trang/màn hình. Chi tiết danh mục giao diện thực tế khách hàng nghiệm thu."
    ws3["A2"].font = Font(name=font_family, size=10, italic=True, color="595959")
    ws3["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws3.row_dimensions[2].height = 22

    # Headers (Row 4)
    headers_ws3 = ["Nhóm màn hình", "Trang / Màn hình", "Mục đích sử dụng", "Core", "Full", "Thành phần UI chính"]
    for c_idx, h in enumerate(headers_ws3, 1):
        cell = ws3.cell(4, c_idx, h)
        cell.font = Font(name=font_family, size=10, bold=True, color=c_white)
        cell.fill = PatternFill(start_color=c_navy, end_color=c_navy, fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if c_idx in [4, 5] else "left", vertical="center")
        cell.border = cell_border
    ws3.row_dimensions[4].height = 26

    # 40 Screens (29 Core, 40 Full)
    screens_data = [
        # Trang Công khai Học viện (Core)
        ("Trang chủ & Tổng quan", "Trang chủ (Homepage)", "Định vị thương hiệu học viện ẩm thực, dẫn học viên xem khóa học và đăng ký tư vấn", "✓", "✓", "Hero carousel; USP đào tạo thực chiến; Danh mục khóa học; Khóa học tiêu biểu; Đội ngũ Master Chef; Testimonials học viên; CTA tư vấn; Footer trust"),
        ("Giới thiệu & Uy tín", "Về DuaxCar Kitchen (/ve-duaxcar)", "Xây dựng niềm tin, câu chuyện gìn giữ ẩm thực Việt và triết lý đào tạo thực chiến mở quán", "✓", "✓", "Lịch sử hình thành; Tầm nhìn & Sứ mệnh; Giới thiệu xưởng bếp; Cam kết truyền nghề; Đội ngũ Nghệ nhân Bàn tay vàng; CTA tham quan bếp"),
        ("Đào tạo & Khóa học", "Danh mục khóa học tổng hợp (/khoa-hoc)", "Học viên duyệt và tra cứu các khóa học theo hình thức và chuyên đề", "✓", "✓", "Bộ lọc Onsite/Online; Lọc chuyên đề món (Món sáng, Ốc & hải sản, Lẩu nướng...); Thanh tìm kiếm; Thẻ khóa học trực quan; Trạng thái tuyển sinh"),
        ("Đào tạo & Khóa học", "Khóa học theo chuyên đề món", "Trang đích chuyên sâu cho từng nhóm món ăn cụ thể", "✓", "✓", "Breadcrumb; Giới thiệu chuyên đề; Bộ lọc thời lượng & học phí; Danh sách khóa học thuộc chuyên đề; Form đăng ký chuyên đề"),
        ("Đào tạo & Khóa học", "Chi tiết khóa học (/khoa-hoc/[slug])", "Trang chuyển đổi ra quyết định đăng ký học quan trọng nhất", "✓", "✓", "Hero khóa học; Gallery ảnh món ăn; Thông số sĩ số & thời lượng; Highlights quyền lợi; Curriculum Accordion từng buổi; Giảng viên đứng lớp; Sticky CTA bar; Related courses"),
        ("Đào tạo & Khóa học", "Lịch khai giảng (/lich-khai-giang)", "Xem lịch mở lớp trong tháng và giữ chỗ lớp phù hợp thời gian", "✓", "✓", "Dòng thời gian các lớp mở trong tháng; Ca học sáng/chiều/tối; Địa điểm xưởng bếp; Trạng thái slot còn chỗ; Nút Đăng ký giữ chỗ nhanh"),
        ("Đào tạo & Khóa học", "Đội ngũ Giảng viên", "Danh sách các Master Chef và chuyên gia ẩm thực hàng đầu", "✓", "✓", "Danh sách Master Chef; Danh hiệu Bàn tay vàng; Số năm kinh nghiệm; Triết lý nghề; Các khóa học phụ trách"),
        ("Đào tạo & Khóa học", "Chi tiết Giảng viên", "Hồ sơ chuyên gia chi tiết gia tăng uy tín học viện", "✓", "✓", "Chân dung chất lượng cao; Tiểu sử sự nghiệp; Bằng cấp & Giải thưởng; Quote tâm đắc; Danh sách khóa giảng dạy trực tiếp"),
        ("Tuyển sinh & Lead", "Modal Đăng ký tư vấn khóa học", "Thu thập thông tin học viên nhanh từ mọi điểm chạm trên website", "✓", "✓", "Form nhập Họ tên, SĐT, Email, Khóa học quan tâm, Cơ sở đào tạo, Nhu cầu học; Nút gửi yêu cầu; Trạng thái loading"),
        ("Tuyển sinh & Lead", "Màn hình Xác nhận Đăng ký thành công", "Xác nhận đã nhận thông tin và hướng dẫn các bước tiếp theo", "✓", "✓", "Thông báo thành công; Mã đơn đăng ký; Lời nhắc chuyên viên sẽ gọi lại; Hotline hỗ trợ nhanh"),
        ("Liên hệ & Tư vấn", "Trang Liên hệ (/lien-he)", "Kênh kết nối trực tiếp, đặt lịch tham quan xưởng bếp", "✓", "✓", "Hotline; Zalo OA; Email tuyển sinh; Địa chỉ các cơ sở đào tạo; Form để lại tin nhắn; Bản đồ Google Maps nhúng"),
        ("Hỗ trợ & Hỏi đáp", "Trang FAQ Hỏi đáp (/faq)", "Giải đáp nhanh các băn khoăn về học phí, lịch học, mở quán", "✓", "✓", "Tìm kiếm câu hỏi; Accordion câu hỏi theo chủ đề (Học phí, Bằng cấp, Hỗ trợ mở quán, Địa điểm); CTA tư vấn trực tiếp"),
        ("Chính sách & Pháp lý", "Quy chế đào tạo & Điều khoản (/chinh-sach/dieu-khoan)", "Quy định học viên, quyền lợi và nghĩa vụ khi tham gia học", "✓", "✓", "Quy chế lớp học; Cam kết chất lượng; Quyền bảo lưu khóa học; Bản quyền công thức"),
        ("Chính sách & Pháp lý", "Chính sách bảo mật thông tin (/chinh-sach/bao-mat)", "Cam kết bảo mật dữ liệu học viên theo quy định pháp luật", "✓", "✓", "Mục đích thu thập; Phạm vi sử dụng; Cam kết không chia sẻ bên thứ ba; Quyền yêu cầu xóa thông tin"),
        ("Chính sách & Pháp lý", "Chính sách học phí & Hoàn cọc (/chinh-sach/thanh-toan)", "Quy chế thu học phí, quy định cọc giữ chỗ và điều kiện hoàn bảo lưu", "✓", "✓", "Hình thức đóng học phí; Quy định cọc giữ chỗ; Thời hạn bảo lưu học tập 12 tháng; Điều kiện giải quyết hoàn phí"),
        ("Nội dung & Cẩm nang", "Danh sách Cẩm nang ẩm thực (/tin-tuc)", "Chia sẻ kiến thức nấu ăn và kinh nghiệm kinh doanh F&B thu hút traffic", "✓", "✓", "Featured article; Danh mục chuyên đề; Thẻ bài viết; Tác giả; Thời gian đọc; Tìm kiếm bài viết"),
        ("Nội dung & Cẩm nang", "Chi tiết bài viết Cẩm nang (/tin-tuc/[slug])", "Bài viết chuẩn SEO cung cấp giá trị chuyên sâu cho người đọc", "✓", "✓", "Tiêu đề chuẩn SEO; Tác giả Master Chef; Mục lục tự động (TOC); Nội dung định dạng đẹp; Khối bài liên quan; CTA khóa học"),
        ("Hệ thống", "Trang 404 Không tìm thấy trang", "Điều hướng học viên khi truy cập link sai hoặc hết hạn", "✓", "✓", "Thông báo thân thiện; Nút quay lại trang chủ; Gợi ý các khóa học đang tuyển sinh nổi bật"),
        
        # Màn hình Quản trị CMS Core
        ("Quản trị CMS", "Đăng nhập Quản trị (/login)", "Xác thực an toàn cho cán bộ tuyển sinh và quản trị viên", "✓", "✓", "Form email/password; Xác thực Supabase Auth; Thông báo lỗi bảo mật"),
        ("Quản trị CMS", "Dashboard Tổng quan Tuyển sinh (/admin)", "Thống kê chỉ số tuyển sinh và công việc cần xử lý trong ngày", "✓", "✓", "Thống kê lead mới; Tổng số khóa học; Danh sách học viên cần gọi tư vấn gấp; Lối tắt chức năng"),
        ("Quản trị CMS", "Quản lý Khóa học (/admin/khoa-hoc)", "Danh sách toàn bộ khóa học trong học viện", "✓", "✓", "Bảng khóa học; Bộ lọc chuyên mục; Bật/tắt trạng thái nổi bật; Nút thêm mới và sửa"),
        ("Quản trị CMS", "Chỉnh sửa Khóa học & Curriculum", "Form chi tiết cấu hình khóa học và giáo trình theo buổi", "✓", "✓", "Thông tin khóa, học phí, sĩ số; Curriculum Builder từng buổi; Highlights tags; Chọn ảnh Media"),
        ("Quản trị CMS", "Quản lý Giảng viên (/admin/giang-vien)", "Danh sách giảng viên Master Chef", "✓", "✓", "Bảng giảng viên; Thêm mới; Chỉnh sửa thành tựu; Gán khóa học; Chọn ảnh chân dung"),
        ("Quản trị CMS", "Quản lý Đơn đăng ký học viên (/admin/dang-ky)", "Quản lý danh sách Lead tuyển sinh và trạng thái tư vấn", "✓", "✓", "Bảng lead; Lọc theo khóa học & trạng thái; Tìm kiếm SĐT; Cập nhật trạng thái xử lý; Ghi chú nội bộ"),
        ("Quản trị CMS", "Quản lý Bài viết Cẩm nang (/admin/tin-tuc)", "Danh sách các bài viết tin tức & cẩm nang ẩm thực", "✓", "✓", "Bảng bài viết; Lọc chuyên mục; Trạng thái hiển thị; Nút thêm mới và chỉnh sửa"),
        ("Quản trị CMS", "Soạn thảo Bài viết Rich Text Editor", "Trình soạn thảo văn bản phong phú cho bài viết blog", "✓", "✓", "Rich Text Editor; Định dạng tiêu đề H2/H3; Chèn ảnh từ Media; Cấu hình SEO slug & excerpt"),
        ("Quản trị CMS", "Thư viện Media Tập trung (/admin/media)", "Quản lý kho tài nguyên hình ảnh món ăn và video", "✓", "✓", "Lưới ảnh thumbnail; Nén ảnh WebP tự động khi upload; Lightbox phóng to; Sao chép URL; Xóa ảnh"),
        ("Quản trị CMS", "Quản lý FAQ & Chính sách (/admin/faq, /admin/chinh-sach)", "Quản lý câu hỏi thường gặp và văn bản pháp lý", "✓", "✓", "CRUD câu hỏi FAQ theo nhóm; Trình biên tập nội dung các trang chính sách pháp lý"),
        ("Quản trị CMS", "Cài đặt Hệ thống & Thương hiệu (/admin/cai-dat)", "Cấu hình hotline, địa chỉ cơ sở, social và SEO mặc định", "✓", "✓", "Cấu hình tên trung tâm, slogan, hotline, địa chỉ các bếp, link fanpage/zalo/tiktok, banner top"),

        # Màn hình Mở rộng Gói Full (11 screens Add-on)
        ("Full – Thanh toán", "Cổng chuyển khoản VietQR tự động", "Học viên quét mã QR thanh toán cọc giữ chỗ tức thì", "—", "✓", "Mã QR động đúng số tiền cọc; Thông tin tài khoản ngân hàng trung tâm; Cú pháp chuyển khoản; Hướng dẫn xác nhận"),
        ("Full – Thanh toán", "Kết quả thanh toán học phí qua Gateway", "Thông báo kết quả sau khi thanh toán qua VNPay/MoMo", "—", "✓", "Mã giao dịch điện tử; Trạng thái thành công/thất bại; Biên lai đặt cọc; Nút quay về trang khóa học"),
        ("Full – Cổng Học viên", "Đăng nhập / Đăng ký Cổng học viên", "Học viên đăng nhập để tra cứu lịch học và tài liệu công thức", "—", "✓", "Đăng nhập bằng SĐT/OTP hoặc Email; Quên mật khẩu; Đổi mật khẩu học viên"),
        ("Full – Cổng Học viên", "Dashboard Cá nhân Học viên (/hoc-vien)", "Khu vực cá nhân theo dõi toàn bộ quá trình học nghề", "—", "✓", "Khóa học của tôi; Lịch khai giảng lớp đã đăng ký; Thông tin phòng bếp thực hành; Trạng thái học phí"),
        ("Full – Cổng Học viên", "Kho Tài liệu Công thức Số Độc quyền", "Học viên tải công thức chuẩn định lượng có Watermark chống lộ", "—", "✓", "Danh sách tài liệu theo khóa học; File PDF công thức chuẩn gram; Video kỹ thuật mẫu; Danh bạ nguồn nguyên liệu sỉ"),
        ("Full – Social Proof", "Showcase Học viên Mở quán Thành công", "Trang trưng bày các quán ăn thực tế của học viên đã thành nghề", "—", "✓", "Hình ảnh quán ăn học viên; Video chia sẻ thực tế; Doanh thu quán; Món bán chạy; Địa chỉ quán ăn"),
        ("Full – Social Proof", "Form Gửi đánh giá & Upload ảnh thành phẩm", "Học viên sau tốt nghiệp gửi cảm nhận và hình ảnh món ăn", "—", "✓", "Form chấm điểm sao (1-5 sao); Nhập cảm nhận khóa học; Tải ảnh món tự nấu; Gửi duyệt lên ban quản trị"),
        ("Full – Tuyển sinh Pro", "Quản lý Lịch Khai giảng & Slot Học viên (CMS)", "Quản lý chi tiết từng lớp khai giảng trong CMS", "—", "✓", "Thêm lớp mới; Thiết lập ngày học, ca học, phòng bếp; Giới hạn sĩ số; Tự động khóa slot khi đủ cọc"),
        ("Full – Tuyển sinh Pro", "Báo cáo Doanh thu Tuyển sinh & Lead Funnel (CMS)", "Phân tích hiệu quả tuyển sinh và doanh thu đào tạo", "—", "✓", "Biểu đồ số lượng lead theo nguồn; Tỷ lệ chuyển đổi thành học viên nộp cọc; Doanh thu học phí theo tháng/quý"),
        ("Full – Vận hành Pro", "Quản lý Giao dịch Học phí & Đối soát (CMS)", "Kế toán quản lý danh sách thu cọc và biên lai học phí", "—", "✓", "Danh sách giao dịch VietQR và Gateway; Trạng thái thanh toán; Xác nhận thu cọc thủ công; Xuất file Excel"),
        ("Full – Vận hành Pro", "Quản lý Phân quyền Nhân sự Đa cấp (CMS)", "Thiết lập quyền truy cập cho từng bộ phận trong trung tâm", "—", "✓", "Quản lý tài khoản cán bộ tuyển sinh, giáo vụ, kế toán; Cấu hình quyền xem/sửa/xóa tương ứng")
    ]

    for r_idx, screen in enumerate(screens_data, 5):
        ws3.row_dimensions[r_idx].height = 22
        group, s_name, s_purp, s_core, s_full, s_elem = screen
        
        ws3.cell(r_idx, 1, group).alignment = Alignment(horizontal="left", vertical="center")
        ws3.cell(r_idx, 1).font = Font(name=font_family, size=9, bold=True, color=c_core_txt)
        
        ws3.cell(r_idx, 2, s_name).alignment = Alignment(horizontal="left", vertical="center")
        ws3.cell(r_idx, 2).font = Font(name=font_family, size=9, bold=True)
        
        ws3.cell(r_idx, 3, s_purp).alignment = Alignment(horizontal="left", vertical="center")
        ws3.cell(r_idx, 3).font = Font(name=font_family, size=9)
        
        c_cell = ws3.cell(r_idx, 4, s_core)
        c_cell.alignment = Alignment(horizontal="center", vertical="center")
        c_cell.font = Font(name=font_family, size=10, bold=(s_core=="✓"), color=c_core_txt if s_core=="✓" else "7F7F7F")
        if s_core == "✓":
            c_cell.fill = PatternFill(start_color=c_check_bg, end_color=c_check_bg, fill_type="solid")

        f_cell = ws3.cell(r_idx, 5, s_full)
        f_cell.alignment = Alignment(horizontal="center", vertical="center")
        f_cell.font = Font(name=font_family, size=10, bold=True, color=c_green_txt)
        if s_full == "✓":
            f_cell.fill = PatternFill(start_color=c_green_bg, end_color=c_green_bg, fill_type="solid")

        ws3.cell(r_idx, 6, s_elem).alignment = Alignment(horizontal="left", vertical="center")
        ws3.cell(r_idx, 6).font = Font(name=font_family, size=9)

        if r_idx % 2 == 0:
            for col_i in [1, 2, 3, 6]:
                ws3.cell(r_idx, col_i).fill = PatternFill(start_color=c_alt_row, end_color=c_alt_row, fill_type="solid")

        for col_i in range(1, 7):
            ws3.cell(r_idx, col_i).border = cell_border

    last_screen_row = len(screens_data) + 4
    screen_summary_row = last_screen_row + 1

    ws3.cell(screen_summary_row, 3, "TỔNG SỐ MÀN HÌNH UI BÀN GIAO").alignment = Alignment(horizontal="right", vertical="center")
    ws3.cell(screen_summary_row, 3).font = Font(name=font_family, size=10, bold=True, color=c_navy)

    ws3.cell(screen_summary_row, 4, f"=COUNTIF(D5:D{last_screen_row},\"✓\")").alignment = Alignment(horizontal="center", vertical="center")
    ws3.cell(screen_summary_row, 4).font = Font(name=font_family, size=11, bold=True, color=c_core_txt)
    ws3.cell(screen_summary_row, 4).fill = PatternFill(start_color=c_gray_bg, end_color=c_gray_bg, fill_type="solid")

    ws3.cell(screen_summary_row, 5, f"=COUNTIF(E5:E{last_screen_row},\"✓\")").alignment = Alignment(horizontal="center", vertical="center")
    ws3.cell(screen_summary_row, 5).font = Font(name=font_family, size=11, bold=True, color=c_green_txt)
    ws3.cell(screen_summary_row, 5).fill = PatternFill(start_color=c_green_bg, end_color=c_green_bg, fill_type="solid")

    for col_i in range(1, 7):
        ws3.cell(screen_summary_row, col_i).border = thick_bottom

    # Link screens count in Sheet 1 to summary row
    ws1["E7"] = f"='03_TRANG_UI'!D{screen_summary_row}"
    ws1["H7"] = f"='03_TRANG_UI'!E{screen_summary_row}"

    widths_ws3 = {'A': 22.0, 'B': 36.0, 'C': 44.0, 'D': 10.0, 'E': 10.0, 'F': 80.0}
    for col_letter, width in widths_ws3.items():
        ws3.column_dimensions[col_letter].width = width

    # =========================================================================
    # SHEET 4: 04_CHI_PHI_DINH_KY (Recurring / 3rd Party Costs)
    # =========================================================================
    ws4 = wb.create_sheet(title="04_CHI_PHI_DINH_KY")
    ws4.views.sheetView[0].showGridLines = True

    # Title
    ws4.merge_cells("A1:G1")
    ws4["A1"] = "CHI PHÍ ĐỊNH KỲ / BÊN THỨ BA – TÁCH KHỎI PHÍ PHÁT TRIỂN PHẦN MỀM"
    ws4["A1"].font = Font(name=font_family, size=14, bold=True, color=c_white)
    ws4["A1"].fill = PatternFill(start_color=c_navy, end_color=c_navy, fill_type="solid")
    ws4["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws4.row_dimensions[1].height = 36

    # Headers (Row 3)
    headers_ws4 = ["Hạng mục", "Đơn vị", "SL", "Đơn giá (nhập)", "Chu kỳ/năm", "Chi phí năm", "Ghi chú"]
    for c_idx, h in enumerate(headers_ws4, 1):
        cell = ws4.cell(3, c_idx, h)
        cell.font = Font(name=font_family, size=10, bold=True, color=c_white)
        cell.fill = PatternFill(start_color=c_navy, end_color=c_navy, fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if c_idx in [2, 3, 5] else "left", vertical="center")
        cell.border = cell_border
    ws4.row_dimensions[3].height = 26

    recurring_items = [
        ("Tên miền thương hiệu (.vn hoặc .com)", "năm", 1, None, 1, "Tùy chọn tên miền thương hiệu: .vn (~550.000đ/năm) hoặc .com (~300.000đ/năm)."),
        ("Cloud Hosting / Server (Vercel Pro / VPS)", "tháng", 1, None, 12, "Lưu trữ máy chủ ứng dụng Next.js; gói Pro (~480.000đ/tháng) hoặc VPS riêng."),
        ("Cơ sở dữ liệu & Storage (Supabase Pro)", "tháng", 1, None, 12, "Database PostgreSQL, Authentication và lưu trữ ảnh/video; gói Pro ~600.000đ/tháng."),
        ("CDN / SSL / WAF / Bot Protection (Cloudflare)", "tháng", 1, None, 12, "Tăng tốc tải trang và chống tấn công DDoS; có thể dùng gói Free hoặc Pro theo nhu cầu."),
        ("Email Transactional (Resend / Brevo)", "tháng", 1, None, 12, "Gửi email xác nhận đăng ký học, lịch khai giảng và biên nhận cọc tự động (miễn phí đến 3.000 email/tháng)."),
        ("SMS Brandname / Zalo ZNS Tuyển sinh", "tháng", 1, None, 12, "Gửi tin nhắn thương hiệu xác nhận giữ chỗ và nhắc lịch học; tính theo số lượng gửi thực tế (~300đ/tin)."),
        ("Cổng thanh toán trực tuyến / VietQR", "giao dịch", 1, None, 1, "Cổng thanh toán VietQR miễn phí; Cổng VNPay/MoMo thu phí ~1.1% - 2.0% trên mỗi giao dịch thành công."),
        ("Bảo trì, Backup & Hỗ trợ kỹ thuật SLA", "tháng", 1, None, 12, "Gói hỗ trợ vận hành định kỳ sau thời gian bảo hành tiêu chuẩn (tùy chọn theo thỏa thuận SLA).")
    ]

    for idx, item in enumerate(recurring_items, 4):
        ws4.row_dimensions[idx].height = 22
        name, unit, qty, unit_price, freq, note = item
        
        ws4.cell(idx, 1, name).alignment = Alignment(horizontal="left", vertical="center")
        ws4.cell(idx, 1).font = Font(name=font_family, size=9, bold=True, color=c_core_txt)
        
        ws4.cell(idx, 2, unit).alignment = Alignment(horizontal="center", vertical="center")
        ws4.cell(idx, 2).font = Font(name=font_family, size=9)
        
        ws4.cell(idx, 3, qty).alignment = Alignment(horizontal="center", vertical="center")
        ws4.cell(idx, 3).font = Font(name=font_family, size=9)
        
        # Unit price input cell
        up_cell = ws4.cell(idx, 4, unit_price)
        up_cell.alignment = Alignment(horizontal="right", vertical="center")
        up_cell.font = Font(name=font_family, size=9)
        up_cell.number_format = '#,##0\\ "₫"'
        up_cell.fill = PatternFill(start_color=c_yellow_accent, end_color=c_yellow_accent, fill_type="solid")
        
        ws4.cell(idx, 5, freq).alignment = Alignment(horizontal="center", vertical="center")
        ws4.cell(idx, 5).font = Font(name=font_family, size=9)
        
        # Formula: =IF(OR(D{idx}="",E{idx}=0),"",C{idx}*D{idx}*E{idx})
        tot_cell = ws4.cell(idx, 6, f"=IF(OR(D{idx}=\"\",E{idx}=0),\"\",C{idx}*D{idx}*E{idx})")
        tot_cell.alignment = Alignment(horizontal="right", vertical="center")
        tot_cell.font = Font(name=font_family, size=9, bold=True)
        tot_cell.number_format = '#,##0\\ "₫"'
        
        ws4.cell(idx, 7, note).alignment = Alignment(horizontal="left", vertical="center")
        ws4.cell(idx, 7).font = Font(name=font_family, size=9, italic=True)

        for col_i in range(1, 8):
            ws4.cell(idx, col_i).border = cell_border

    # Total row
    sum_row_ws4 = len(recurring_items) + 5
    ws4.cell(sum_row_ws4, 5, "TỔNG CHI PHÍ NĂM").alignment = Alignment(horizontal="right", vertical="center")
    ws4.cell(sum_row_ws4, 5).font = Font(name=font_family, size=10, bold=True, color=c_navy)

    total_rec_cell = ws4.cell(sum_row_ws4, 6, f"=SUM(F4:F{sum_row_ws4-2})")
    total_rec_cell.alignment = Alignment(horizontal="right", vertical="center")
    total_rec_cell.font = Font(name=font_family, size=11, bold=True, color=c_core_txt)
    total_rec_cell.number_format = '#,##0\\ "₫"'
    total_rec_cell.fill = PatternFill(start_color=c_gray_bg, end_color=c_gray_bg, fill_type="solid")

    for col_i in range(1, 8):
        ws4.cell(sum_row_ws4, col_i).border = thick_bottom

    widths_ws4 = {'A': 36.0, 'B': 14.0, 'C': 8.0, 'D': 18.0, 'E': 14.0, 'F': 22.0, 'G': 65.0}
    for col_letter, width in widths_ws4.items():
        ws4.column_dimensions[col_letter].width = width

    # =========================================================================
    # SHEET 5: 05_PHAM_VI_LUU_Y (Scoping Parameters & Conditions)
    # =========================================================================
    ws5 = wb.create_sheet(title="05_PHAM_VI_LUU_Y")
    ws5.views.sheetView[0].showGridLines = True

    # Title
    ws5.merge_cells("A1:F1")
    ws5["A1"] = "PHẠM VI & ĐIỀU KIỆN KỸ THUẬT CẦN CHỐT TRƯỚC KHI TRIỂN KHAI"
    ws5["A1"].font = Font(name=font_family, size=14, bold=True, color=c_white)
    ws5["A1"].fill = PatternFill(start_color=c_navy, end_color=c_navy, fill_type="solid")
    ws5["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws5.row_dimensions[1].height = 36

    # Headers (Row 3)
    headers_ws5 = ["STT", "Hạng mục", "Cần chốt cụ thể", "Ví dụ minh họa thực tế", "Gói Tiết kiệm (Core)", "Gói Full (Khuyến nghị)"]
    for c_idx, h in enumerate(headers_ws5, 1):
        cell = ws5.cell(3, c_idx, h)
        cell.font = Font(name=font_family, size=10, bold=True, color=c_white)
        cell.fill = PatternFill(start_color=c_navy, end_color=c_navy, fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if c_idx == 1 else "left", vertical="center")
        cell.border = cell_border
    ws5.row_dimensions[3].height = 26

    scope_items = [
        (1, "Danh mục Khóa học & Số lượng", "Số lượng khóa học ban đầu đưa lên website; cấu trúc phân loại chuyên đề; thông tin chi tiết từng khóa.", "Ví dụ: 15 khóa học ban đầu (Món sáng, Lẩu nướng, Món ốc, Bếp trưởng khởi nghiệp); trường dữ liệu: giá, số buổi, sĩ số, highlights, giáo trình.", "Nhập sẵn 10-15 khóa học mẫu chuẩn cấu trúc; hướng dẫn admin tự tạo không giới hạn.", "Hỗ trợ import toàn bộ khóa học, chuẩn hóa hình ảnh món và cấu trúc giáo trình đa cấp."),
        (2, "Giáo trình Đào tạo (Curriculum)", "Mức độ chi tiết của giáo trình khóa học; số buổi học; nội dung công thức và định lượng nguyên liệu.", "Ví dụ: Mỗi khóa có 3-5 buổi học; mỗi buổi gồm tiêu đề, tóm tắt kỹ năng đạt được và chi tiết các công thức chế biến.", "Giáo trình interactive accordion theo buổi học; quản lý linh hoạt qua Curriculum JSON editor.", "Tích hợp file giáo trình số PDF, video clip quay mẫu và tài liệu bảo mật Watermark."),
        (3, "Hồ sơ Giảng viên & Master Chef", "Số lượng giảng viên đưa lên hệ thống; thông tin bằng cấp, danh hiệu và giải thưởng ẩm thực.", "Ví dụ: 3-5 Master Chef; Thầy Phạm Văn Long (Nghệ nhân Bàn tay vàng, 25+ năm kinh nghiệm); ảnh chân dung, giải thưởng, khóa phụ trách.", "Tạo hồ sơ chuyên gia chuẩn nhận diện uy tín, gán khóa học phụ trách tương ứng.", "Thêm trang chi tiết chuyên sâu từng Master Chef, chứng chỉ nghề và bài phỏng vấn truyền nghề."),
        (4, "Lịch Khai giảng & Quản lý Slot", "Quy định mở lớp hàng tháng; số cơ sở đào tạo; quy tắc hiển thị trạng thái lớp học (Còn chỗ / Sắp hết / Đã đóng).", "Ví dụ: Mỗi tháng mở 6-8 lớp; ca học Sáng (8h30-11h30), Chiều (14h-17h); giới hạn tối đa 8-10 học viên/lớp để cầm tay chỉ việc.", "Hiển thị lịch khai giảng theo tháng; đăng ký tư vấn chọn lớp thủ công.", "Tự động khóa slot khi nhận đủ số lượng cọc; hiển thị số chỗ còn lại theo thời gian thực."),
        (5, "Quy trình Thu thập Lead & Tuyển sinh", "Các trường thông tin học viên cần thu thập; quy trình tiếp nhận và phân bổ chuyên viên tư vấn.", "Ví dụ: Họ tên, Số điện thoại (bắt buộc), Email, Khóa quan tâm, Cơ sở đào tạo mong muốn, Nhu cầu (mở quán / nâng tay nghề).", "Thu lead chuẩn vào database Supabase; quản lý trạng thái tư vấn trong CMS back-office.", "Tự động gửi email/ZNS xác nhận; đồng bộ dữ liệu sang Google Sheets/CRM ngoài qua Webhook."),
        (6, "Thiết kế Giao diện (UI/UX)", "Số lượng màn hình thiết kế; phong cách nhận diện thương hiệu; số vòng chỉnh sửa góp ý.", "Ví dụ: Nhận diện màu Cam-Đen DuaxCar Kitchen; phong cách hiện đại, ngon miệng, ẩm thực cao cấp; tối đa 2 vòng góp ý hoàn thiện.", "Thiết kế hoàn chỉnh 29 màn hình Core desktop & mobile theo nhận diện DuaxCar.", "Bao gồm 40 màn hình toàn diện (thêm Cổng học viên, Showcase quán học viên, Cổng cọc VietQR)."),
        (7, "Trang Nội dung & Cẩm nang Ẩm thực", "Danh sách các trang tĩnh cần xây dựng và số bài viết cẩm nang ban đầu khi bàn giao website.", "Ví dụ: Giới thiệu, Liên hệ, FAQ, Quy chế đào tạo, Bảo mật, Học phí & bảo lưu; nhập sẵn 10 bài viết cẩm nang ẩm thực chuẩn SEO.", "Các trang tĩnh cốt lõi + 10 bài viết cẩm nang nhập sẵn ban đầu.", "Mở rộng bài viết cẩm nang, bí quyết mở quán và mục lục tự động (Table of Contents)."),
        (8, "Cổng Thanh toán Đặt cọc Giữ chỗ", "Phương thức thanh toán cọc giữ chỗ; tài khoản ngân hàng nhận tiền; quy trình đối soát.", "Ví dụ: Sinh mã VietQR động đúng số tiền cọc (1.000.000đ) kèm mã đơn; hoặc cổng VNPay/MoMo liên kết tài khoản doanh nghiệp.", "Hướng dẫn chuyển khoản ngân hàng thủ công kèm cú pháp mã đơn.", "Tự động sinh mã VietQR động theo đơn; tích hợp cổng thanh toán trực tuyến và Webhook đối soát."),
        (9, "Social Proof & Học viên Mở quán", "Cơ chế thu thập và kiểm duyệt cảm nhận học viên; hình ảnh quán ăn thực tế sau khóa học.", "Ví dụ: Hình ảnh học viên mở quán phở/lẩu tại Hà Nội & các tỉnh; video phỏng vấn doanh thu quán; điểm đánh giá 5 sao.", "Khu vực Testimonials chọn lọc hiển thị trang chủ.", "Showcase chuyên đề học viên mở quán thành công; form gửi review xác thực kèm ảnh thành phẩm."),
        (10, "Cổng Học viên & Bảo vệ Bản quyền Công thức", "Quy định tài khoản học viên; bảo mật tài liệu công thức bí quyết ẩm thực độc quyền của DuaxCar.", "Ví dụ: Học viên đăng nhập xem tài liệu PDF công thức; tự động đóng dấu chìm Watermark SĐT học viên lên file.", "Cung cấp tài liệu giáo trình trực tiếp tại xưởng bếp trong buổi học.", "Cổng học viên riêng biệt; kho công thức điện tử có Watermark chống sao chép và phát tán."),
        (11, "Thông báo Tự động (Email & ZNS)", "Kênh gửi thông báo và các sự kiện tự động kích hoạt thông báo đến học viên.", "Ví dụ: Gửi email chào mừng khi đăng ký; tin nhắn Zalo ZNS xác nhận giữ chỗ thành công; nhắc lịch học trước 2 ngày.", "Email thông báo tiếp nhận thông tin học viên cơ bản.", "Hệ thống Email tự động theo kịch bản tuyển sinh + Tích hợp Zalo ZNS / SMS Brandname."),
        (12, "SEO & Tiếp thị Số (Tracking)", "Bộ từ khóa mục tiêu ngành đào tạo nghề F&B; mã theo dõi quảng cáo cần tích hợp.", "Ví dụ: Từ khóa 'học nấu ăn mở quán', 'học làm nước dùng phở'; mã GA4, Google Tag Manager, Meta Pixel theo dõi form.", "SEO On-page cơ bản, Course Schema chuẩn Google.", "SEO Schema đa tầng, sitemap tự động, thiết lập tracking chuyển đổi toàn diện cho GA4 và Facebook Ads."),
        (13, "Phân quyền Quản trị (CMS Roles)", "Số lượng tài khoản nhân sự quản trị và phạm vi phân quyền của từng nhóm.", "Ví dụ: Ban Giám đốc (toàn quyền); Cán bộ Tuyển sinh (chỉ xem lead); Bếp trưởng (sửa giáo trình); Kế toán (đối soát học phí).", "Tài khoản Admin quản trị chung được bảo vệ qua Supabase Auth.", "Phân quyền đa cấp chi tiết theo phòng ban (Tuyển sinh, Giáo vụ, Kế toán) kèm Audit Log."),
        (14, "Hạ tầng Hosting & Cơ sở dữ liệu", "Nền tảng máy chủ; dung lượng lưu trữ; quy mô lưu lượng truy cập dự kiến; cơ chế sao lưu.", "Ví dụ: Nền tảng Vercel Serverless; database Supabase PostgreSQL; backup tự động hàng ngày; cơ chế Cron Ping giữ DB.", "Cấu hình chuẩn đám mây tối ưu chi phí vận hành cho giai đoạn ra mắt ban đầu.", "Hạ tầng chịu tải cao, CDN tăng tốc toàn cầu, tối ưu hóa bộ nhớ cache ISR và sao lưu tự động."),
        (15, "Bảo mật & An toàn Dữ liệu", "Yêu cầu an toàn thông tin dữ liệu khách hàng; mã hóa và phòng chống tấn công mạng.", "Ví dụ: Row Level Security (RLS) bảo vệ bảng học viên; mã hóa mật khẩu; chống spam form; chứng chỉ SSL/TLS.", "Bảo mật baseline tiêu chuẩn: HTTPS, RLS, bảo vệ biến môi trường.", "Bảo vệ đa lớp: RLS nâng cao, Cloudflare WAF, chống brute-force và kiểm soát phiên đăng nhập."),
        (16, "Bảo hành & Hỗ trợ Kỹ thuật", "Thời gian bảo hành; thời gian phản hồi sự cố; phạm vi khắc phục lỗi phần mềm.", "Ví dụ: Bảo hành kỹ thuật 03 tháng sau nghiệm thu; phản hồi lỗi nghiêm trọng trong vòng 2-4 giờ làm việc.", "Cam kết bảo hành lỗi kỹ thuật 03 tháng kể từ ngày bàn giao nghiệm thu.", "Bảo hành kỹ thuật 06 tháng; ưu đãi giảm 20% khi nâng cấp thêm tính năng trong năm đầu tiên.")
    ]

    for idx, row_data in enumerate(scope_items, 4):
        ws5.row_dimensions[idx].height = 28
        stt, cat, req, eg, core_p, full_p = row_data
        
        ws5.cell(idx, 1, stt).alignment = Alignment(horizontal="center", vertical="center")
        ws5.cell(idx, 1).font = Font(name=font_family, size=9)
        
        ws5.cell(idx, 2, cat).alignment = Alignment(horizontal="left", vertical="center")
        ws5.cell(idx, 2).font = Font(name=font_family, size=9, bold=True, color=c_core_txt)
        
        ws5.cell(idx, 3, req).alignment = Alignment(horizontal="left", vertical="center")
        ws5.cell(idx, 3).font = Font(name=font_family, size=9)
        
        ws5.cell(idx, 4, eg).alignment = Alignment(horizontal="left", vertical="center")
        ws5.cell(idx, 4).font = Font(name=font_family, size=9, italic=True)
        
        ws5.cell(idx, 5, core_p).alignment = Alignment(horizontal="left", vertical="center")
        ws5.cell(idx, 5).font = Font(name=font_family, size=9)
        
        ws5.cell(idx, 6, full_p).alignment = Alignment(horizontal="left", vertical="center")
        ws5.cell(idx, 6).font = Font(name=font_family, size=9)

        if idx % 2 == 0:
            for col_i in range(1, 7):
                ws5.cell(idx, col_i).fill = PatternFill(start_color=c_alt_row, end_color=c_alt_row, fill_type="solid")

        for col_i in range(1, 7):
            ws5.cell(idx, col_i).border = cell_border

    widths_ws5 = {'A': 6.0, 'B': 26.0, 'C': 48.0, 'D': 54.0, 'E': 38.0, 'F': 42.0}
    for col_letter, width in widths_ws5.items():
        ws5.column_dimensions[col_letter].width = width

    # Save to both target locations
    target_path_tulie = "/Users/tungnguyen/Tulie/SeaWay/Bao_gia_Website_DuaxCar_Kitchen_v1_Pham_vi_cu_the.xlsx"
    target_path_codebase = "/Users/tungnguyen/Code/duaxcar/Bao_gia_Website_DuaxCar_Kitchen_v1_Pham_vi_cu_the.xlsx"

    wb.save(target_path_tulie)
    wb.save(target_path_codebase)
    print(f"Successfully generated quotation file:")
    print(f"1. {target_path_tulie}")
    print(f"2. {target_path_codebase}")
    print(f"Total features: {len(raw_features)} (Core: {total_core_items}, Full: {total_full_items})")
    print(f"Total screens: {len(screens_data)}")

if __name__ == "__main__":
    build_duaxcar_quotation()
